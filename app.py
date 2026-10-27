"""
Öğrenci Dekanlığı Kulüp Paneli
--------------------------------
Altınbaş Üniversitesi Öğrenci Dekanlığı'nın kulüpleri, üyeleri ve kulüp
evraklarını yönettiği iç kullanım paneli. Flask ile çalışır, tüm veriyi
data/*.json dosyalarında tutar (bkz. veri.py).

Nasıl çalıştırılır (VS Code terminalinde):
    1) pip install -r requirements.txt
    2) python app.py
    3) Tarayıcıda http://127.0.0.1:5000 adresini açın

Ortam değişkenleri (isteğe bağlı):
    PANEL_SECRET_KEY     -> oturum şifreleme anahtarı
    PANEL_KURTARMA_KODU  -> "Şifremi unuttum" sayfasında istenen kurtarma kodu.
                            Tanımlı değilse şifre sıfırlama kapalıdır.
"""

import hmac
import io
import json
import os
import re
import secrets
from urllib.parse import quote
from datetime import date, datetime, timedelta
from functools import wraps
from uuid import uuid4

from flask import (
    Flask, render_template, abort, request, redirect, url_for, session, flash,
    send_file, send_from_directory,
)
from werkzeug.datastructures import MultiDict
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

import config
import veri

app = Flask(__name__)
app.secret_key = os.environ.get("PANEL_SECRET_KEY", "prototip-gizli-anahtar-degistirin")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # yükleme sınırı: 20 MB

BELGE_KLASORU = os.path.join(app.root_path, "uploads", "belgeler")
BASVURU_KLASORU = os.path.join(app.root_path, "uploads", "basvurular")
ICE_AKTARIM_KLASORU = os.path.join(app.root_path, "uploads", "ice_aktarim")
LOGO_KLASORU = os.path.join(app.static_folder, "images")

BELGE_UZANTILARI = {"pdf", "doc", "docx", "odt", "xls", "xlsx", "ods", "png", "jpg", "jpeg"}
LOGO_UZANTILARI = {"png", "jpg", "jpeg", "webp"}
BELGE_TURLERI = ["Tüzük", "Faaliyet Raporu", "Faaliyet Bildirim Formu", "Faaliyet Planı",
                 "Yönetim Kurulu Listesi", "Danışman Onay Formu", "Danışman Formu", "Etkinlik Başvurusu",
                 "Toplantı Tutanağı", "Diğer"]
YK_GOREVLERI = ["Başkan", "Başkan Yardımcısı", "Sekreter", "Sayman"]
EPOSTA_DESENI = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")

# Giriş yapmadan açılabilen sayfalar.
ACIK_SAYFALAR = {"giris_yap", "sifremi_unuttum", "static"}

# Yalnızca dekanın açabildiği sayfalar: silme, arşiv yönetimi, kullanıcı ve ayarlar.
# Dekanlık personeli bunların dışındaki her şeyi görüntüleyip ekleyip düzenleyebilir.
DEKAN_SAYFALARI = {
    "kulup_sil", "uye_sil", "belge_sil", "danisman_sil", "donem_sil", "basvuru_sil", "zorunlu_evrak_sil",
    "arsiv_geri_getir", "arsiv_kalici_sil",
    "kullanicilar", "kullanici_form", "kullanici_durum", "kullanici_sifre_sifirla", "ayarlar",
}

# Kulüp başkanının açabildiği sayfalar; hepsinde yalnızca kendi kulübünün kayıtlarını görür.
BASKAN_SAYFALARI = {
    "ana_panel", "kulup_detay", "belgeler", "belge_indir", "etkinlik_basvurulari", "basvuru_form",
    "basvuru_detay", "basvuru_dosya", "hesabim", "cikis_yap",
}


def _rol():
    return session.get("rol", "")


def izin(endpoint):
    """Giriş yapmış kullanıcının rolü bu sayfayı açmaya yetiyor mu? (sahiplik kontrolü hariç)"""
    rol = _rol()
    if rol == "dekan":
        return True
    if rol == "personel":
        return endpoint not in DEKAN_SAYFALARI
    if rol == "baskan":
        return endpoint in BASKAN_SAYFALARI
    return False


def _baskan_kulubu():
    """Kulüp başkanının kulüp id'si; başkan değilse None."""
    return session.get("kulup_id") if _rol() == "baskan" else None


def _baskan_sahipligi_kontrol():
    """Kulüp başkanı adresle başka bir kulübün kaydına ulaşmaya çalışırsa 403."""
    kulup_id = _baskan_kulubu()
    argumanlar = request.view_args or {}
    if "kulup_id" in argumanlar and argumanlar["kulup_id"] != kulup_id:
        abort(403)
    if "belge_id" in argumanlar:
        belge = veri.belge_getir(argumanlar["belge_id"])
        if belge is None or belge["kulup_id"] != kulup_id:
            abort(403)
    if "basvuru_id" in argumanlar:
        basvuru = veri.basvuru_getir(argumanlar["basvuru_id"])
        if basvuru is None or basvuru["kulup_id"] != kulup_id:
            abort(403)


def _oturumu_kapat():
    # Sadece kullanıcı bilgisi silinir; csrf anahtarı kalmalı, yoksa
    # tarayıcının arka planda istediği /favicon.ico gibi adresler açık
    # giriş formunun anahtarını geçersiz kılar.
    for anahtar in ("kullanici_adi", "ad_soyad", "rol", "kulup_id", "son_islem"):
        session.pop(anahtar, None)


# ---------------------------------------------------------------------------
# Oturum ve güvenlik
# ---------------------------------------------------------------------------

@app.before_request
def oturum_kontrolu():
    """Giriş, zaman aşımı ve rol kontrolü.

    Giriş yapılmamışsa, kullanıcı pasif yapıldıysa veya belirli süre işlem
    yapılmadıysa giriş sayfasına yönlendirir; rolün açamayacağı sayfada 403 verir.
    """
    if request.endpoint in ACIK_SAYFALAR:
        return None
    sonra = request.full_path.rstrip("?") if request.method == "GET" else None
    kullanici = veri.kullanici_bul(session.get("kullanici_adi", ""))
    if not veri.panel_yetkilisi_mi(kullanici):
        _oturumu_kapat()
        return redirect(url_for("giris_yap", sonra=sonra))

    simdi = datetime.now().timestamp()
    sure = veri.ayarlari_getir()["oturum_suresi_dk"] * 60
    if simdi - session.get("son_islem", simdi) > sure:
        _oturumu_kapat()
        flash("Uzun süre işlem yapılmadığı için oturumunuz kapatıldı. Lütfen tekrar giriş yapın.", "hata")
        return redirect(url_for("giris_yap", sonra=sonra))
    session["son_islem"] = simdi

    # Rol panelde değiştirildiyse oturumdaki bilgi de güncellenir.
    session["rol"] = veri.kullanici_rolu(kullanici)
    session["kulup_id"] = kullanici.get("kulup_id", "")
    if request.endpoint is not None and not izin(request.endpoint):
        abort(403)
    if _rol() == "baskan":
        _baskan_sahipligi_kontrol()
    return None


@app.before_request
def csrf_kontrolu():
    """Tüm POST isteklerinde formdaki gizli anahtarı oturumdakiyle karşılaştırır."""
    if request.method == "POST":
        beklenen = session.get("csrf")
        gelen = request.form.get("csrf", "")
        if not beklenen or not hmac.compare_digest(beklenen, gelen):
            abort(400)


@app.context_processor
def sablon_degiskenleri():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    giris_var = bool(session.get("kullanici_adi"))
    return {
        "csrf": session["csrf"],
        "aktif_sayfa": request.endpoint or "",
        "aktif_donem": veri.aktif_donem() if giris_var else None,
        "bekleyen_basvuru": veri.bekleyen_basvuru_sayisi() if giris_var and izin("basvuru_karar") else 0,
        "izin": izin,
        "rol": _rol(),
        "rol_ad": veri.ROLLER.get(_rol(), ""),
        "baskan_kulubu": _baskan_kulubu(),
        "oturum_suresi_dk": veri.ayarlari_getir()["oturum_suresi_dk"] if giris_var else 0,
    }


# ---------------------------------------------------------------------------
# İşlem geçmişi
# ---------------------------------------------------------------------------

KAYIT_TURLERI = {
    "kulup": "Kulüp", "uye": "Üye", "danisman": "Danışman", "donem": "Dönem", "belge": "Belge",
    "basvuru": "Etkinlik başvurusu", "zorunlu_evrak": "Zorunlu evrak", "karne": "Karne",
    "kullanici": "Kullanıcı", "ayar": "Ayar", "duyuru": "Duyuru",
}

ALAN_ADLARI = {
    "isim": "Ad", "ad": "Ad", "ad_soyad": "Ad soyad", "kategori": "Kategori", "aciklama": "Açıklama",
    "baskan": "Başkan", "durum": "Durum", "danisman_id": "Danışman", "yonetim_kurulu": "Yönetim kurulu",
    "logo": "Logo", "ogrenci_no": "Öğrenci no", "bolum": "Bölüm", "sinif": "Sınıf", "telefon": "Telefon",
    "eposta": "E-posta", "katilim_tarihi": "Katılım tarihi", "unvan": "Unvan", "fakulte": "Fakülte",
    "baslangic": "Başlangıç", "bitis": "Bitiş", "tur": "Tür", "teslim_tarihi": "Son teslim",
    "tarih_saat": "Tarih-saat", "yer": "Yer", "katilimci": "Katılımcı", "konusmaci_var": "Konuşmacı var",
    "konusmaci": "Konuşmacı", "butce": "Bütçe", "butce_aciklama": "Bütçe açıklaması", "puan": "Puan",
    "not": "Not", "rol": "Rol", "kulup_id": "Kulüp", "aktif": "Aktif", "donem_id": "Dönem",
    "baskan_girisi": "Kulüp başkanı girişi", "oturum_suresi_dk": "Oturum süresi (dk)",
}


def _deger_metni(alan, deger):
    if deger is None or deger == "":
        return ""
    if isinstance(deger, bool):
        return "Evet" if deger else "Hayır"
    if alan == "danisman_id":
        danisman = veri.danisman_getir(deger)
        return danisman["tam_ad"] if danisman else deger
    if alan == "kulup_id":
        kulup = next((k for k in veri.kulupleri_getir(arsiv=None) if k["id"] == deger), None)
        return kulup["isim"] if kulup else deger
    if alan == "donem_id":
        donem = veri.donem_getir(deger)
        return donem["ad"] if donem else deger
    if alan == "rol":
        return veri.ROLLER.get(deger, deger)
    if alan == "yonetim_kurulu":
        return "; ".join(f"{u['gorev']}: {u['isim']}" for u in deger)
    return str(deger)


def _degisiklikler(eski, yeni):
    """İki kaydın, ALAN_ADLARI'ndaki alanlarda farkını [{alan, eski, yeni}] olarak döndürür.

    Yeni kayıtta bulunmayan alanlar (ör. logo yüklenmeden kaydedilen formda "logo")
    değişmemiş sayılır.
    """
    eski, yeni = eski or {}, yeni or {}
    sonuc = []
    for alan, ad in ALAN_ADLARI.items():
        if alan not in (yeni or eski):
            continue
        e, y = _deger_metni(alan, eski.get(alan)), _deger_metni(alan, yeni.get(alan))
        if e != y:
            sonuc.append({"alan": ad, "eski": e, "yeni": y})
    return sonuc


def _islem(islem, kayit_turu, kayit_ad, kayit_id="", kulup_id="", eski=None, yeni=None, aciklama=""):
    """İşlem geçmişine kayıt düşer. Güncellemede değişen alan yoksa kayıt düşülmez."""
    degisiklikler = _degisiklikler(eski, yeni) if (eski is not None or yeni is not None) else []
    if islem == "Güncelleme" and not degisiklikler and not aciklama:
        return
    veri.islem_ekle({
        "kullanici": session.get("ad_soyad", ""),
        "kullanici_adi": session.get("kullanici_adi", ""),
        "rol": _rol(),
        "islem": islem,
        "kayit_turu": kayit_turu,
        "kayit_id": kayit_id,
        "kayit_ad": kayit_ad,
        "kulup_id": kulup_id,
        "degisiklikler": degisiklikler,
        "aciklama": aciklama,
    })


def _guvenli_yonlendirme(hedef):
    """Sadece site içi göreli adreslere yönlendirmeye izin verir."""
    if hedef and hedef.startswith("/") and not hedef.startswith("//"):
        return hedef
    return url_for("ana_panel")


@app.errorhandler(400)
@app.errorhandler(403)
@app.errorhandler(404)
@app.errorhandler(413)
@app.errorhandler(500)
def hata_sayfasi(hata):
    kod = getattr(hata, "code", 500) or 500
    mesajlar = {
        400: "İstek geçersiz veya oturum süresi dolmuş. Sayfayı yenileyip tekrar deneyin.",
        403: "Bu işlem için yetkiniz yok.",
        404: "Aradığınız sayfa bulunamadı.",
        413: "Yüklenen dosya çok büyük (en fazla 20 MB).",
        500: "Beklenmeyen bir hata oluştu.",
    }
    return render_template("eror.html", kod=kod, mesaj=mesajlar.get(kod, "Bir hata oluştu.")), kod


# ---------------------------------------------------------------------------
# Giriş / çıkış
# ---------------------------------------------------------------------------

def _sifre_dogru(kullanici, sifre):
    """Şifreyi doğrular. Eski kayıtlarda şifre düz metin tutulmuşsa, doğru girildiğinde
    hemen werkzeug ile hash'e çevrilir; böylece mevcut kullanıcıların girişi bozulmaz."""
    kayitli = kullanici.get("sifre_hash", "")
    if kayitli.startswith(("scrypt:", "pbkdf2:")):
        return check_password_hash(kayitli, sifre)
    if kayitli and hmac.compare_digest(kayitli, sifre):
        veri.kullanici_kaydet(kullanici["kullanici_adi"], {"sifre_hash": generate_password_hash(sifre)})
        return True
    return False


@app.route("/giris", methods=["GET", "POST"])
def giris_yap():
    hata = None
    if request.method == "POST":
        kullanici_adi = request.form.get("kullanici_adi", "").strip()
        sifre = request.form.get("sifre", "")
        kullanici = veri.kullanici_bul(kullanici_adi)

        if veri.panel_yetkilisi_mi(kullanici) and _sifre_dogru(kullanici, sifre):
            csrf = session.get("csrf")
            session.clear()
            session["csrf"] = csrf or secrets.token_hex(16)
            session["kullanici_adi"] = kullanici["kullanici_adi"]
            session["ad_soyad"] = kullanici.get("ad_soyad") or kullanici["kullanici_adi"]
            session["rol"] = veri.kullanici_rolu(kullanici)
            session["kulup_id"] = kullanici.get("kulup_id", "")
            session["son_islem"] = datetime.now().timestamp()
            return redirect(_guvenli_yonlendirme(request.args.get("sonra")))
        hata = "Kullanıcı adı veya şifre hatalı ya da panele erişim yetkiniz yok."

    return render_template("giris.html", hata=hata)


@app.route("/sifremi-unuttum", methods=["GET", "POST"])
def sifremi_unuttum():
    """Yetkili kullanıcının, kurtarma koduyla yeni şifre belirlemesi."""
    kurtarma_kodu = os.environ.get("PANEL_KURTARMA_KODU", "")
    hata = None
    if request.method == "POST" and kurtarma_kodu:
        kullanici_adi = request.form.get("kullanici_adi", "").strip()
        kod = request.form.get("kurtarma_kodu", "")
        yeni_sifre = request.form.get("yeni_sifre", "")
        yeni_sifre_tekrar = request.form.get("yeni_sifre_tekrar", "")
        kullanici = veri.kullanici_bul(kullanici_adi)

        if not veri.panel_yetkilisi_mi(kullanici) or not hmac.compare_digest(kod, kurtarma_kodu):
            hata = "Kullanıcı adı veya kurtarma kodu hatalı."
        elif len(yeni_sifre) < 8:
            hata = "Yeni şifre en az 8 karakter olmalı."
        elif yeni_sifre != yeni_sifre_tekrar:
            hata = "Girdiğiniz şifreler birbiriyle uyuşmuyor."
        else:
            veri.kullanici_sifre_guncelle(kullanici["kullanici_adi"], generate_password_hash(yeni_sifre))
            flash("Şifreniz güncellendi, yeni şifrenizle giriş yapabilirsiniz.", "basarili")
            return redirect(url_for("giris_yap"))

    return render_template("sifremi_unuttum.html", hata=hata, aktif=bool(kurtarma_kodu))


@app.route("/cikis")
def cikis_yap():
    session.clear()
    flash("Oturumunuz kapatıldı.", "basarili")
    return redirect(url_for("giris_yap"))


# ---------------------------------------------------------------------------
# Ana panel
# ---------------------------------------------------------------------------

@app.route("/")
def ana_panel():
    # Kulüp başkanının ana sayfası kendi kulübüdür.
    if _rol() == "baskan":
        return redirect(url_for("kulup_detay", kulup_id=_baskan_kulubu()))
    kulupler = veri.kulupleri_getir()
    tum_uyeler = veri.tum_uyeleri_getir()
    ozet = {
        "kulup": len(kulupler),
        "uye": len({u["ogrenci_no"] for u in tum_uyeler}),
        "uyelik": len(tum_uyeler),
        "aktif": sum(1 for k in kulupler if k["durum"] == "aktif"),
        "pasif": sum(1 for k in kulupler if k["durum"] != "aktif"),
    }
    en_buyukler = sorted(kulupler, key=lambda k: k["uye_sayisi"], reverse=True)[:6]
    # En yakın tarihli olan üstte.
    bekleyenler = sorted(veri.basvurulari_getir(durum="beklemede"), key=lambda b: b["tarih_saat"])

    # Aktif dönemde süresi geçmiş (eksik) veya yaklaşan evrakı olan aktif kulüpler.
    aktif_donem = veri.aktif_donem()
    evrak_uyarilari = []
    donem_ozeti = None
    if aktif_donem:
        aktifler = [k for k in kulupler if k["durum"] == "aktif"]
        karneler = veri.donem_karneleri(aktif_donem["id"], aktifler)
        puanlar = [k["puan"] for k in karneler if k["puan"] is not None]
        beklenen = sum(k["evrak_toplam"] for k in karneler)
        donem_ozeti = {
            "etkinlik": sum(k["etkinlik"] for k in karneler),
            "teslim_orani": round(100 * sum(k["evrak_teslim"] for k in karneler) / beklenen) if beklenen else None,
            "ortalama_puan": f"{sum(puanlar) / len(puanlar):.1f}".replace(".", ",") if puanlar else None,
            "degerlendirilen": len(puanlar),
            "aktif": len(aktifler),
        }
        _, satirlar = veri.evrak_durumlari(aktif_donem["id"], kulupler=aktifler)
        evrak_uyarilari = sorted(
            (s for s in satirlar if s["sayilar"]["eksik"] or s["sayilar"]["yaklasiyor"]),
            key=lambda s: (-s["sayilar"]["eksik"], -s["sayilar"]["yaklasiyor"], s["kulup"]["isim"]),
        )
    return render_template(
        "panel.html",
        ozet=ozet,
        bekleyenler=bekleyenler,
        evrak_uyarilari=evrak_uyarilari,
        donem_ozeti=donem_ozeti,
        eksikli_kulup=sum(1 for s in evrak_uyarilari if s["sayilar"]["eksik"]),
        son_uyeler=veri.son_eklenen_uyeler(),
        en_buyukler=en_buyukler,
        en_fazla=max((k["uye_sayisi"] for k in en_buyukler), default=0),
    )


# ---------------------------------------------------------------------------
# Kulüpler
# ---------------------------------------------------------------------------

def _kulup_veya_404(kulup_id):
    kulup = veri.kulup_getir(kulup_id)
    if kulup is None:
        abort(404)
    return kulup


_TR_HARF = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")


def _kulup_id_uret(isim):
    temel = re.sub(r"[^a-z0-9]+", "-", isim.translate(_TR_HARF).lower()).strip("-")
    temel = re.sub(r"-?kulubu?$", "", temel) or "kulup"
    mevcutlar = veri.tum_kulup_idleri()  # arşivdekiler dahil
    aday, sayac = temel, 2
    while aday in mevcutlar:
        aday, sayac = f"{temel}-{sayac}", sayac + 1
    return aday


def _uzanti(dosya_adi):
    return dosya_adi.rsplit(".", 1)[-1].lower() if "." in dosya_adi else ""


@app.route("/kulupler")
def kulupler():
    return render_template(
        "kulupler.html",
        kulupler=veri.kulupleri_getir(),
        kategoriler=veri.kategorileri_getir(),
    )


@app.route("/kulupler/yeni", methods=["GET", "POST"])
@app.route("/kulupler/<kulup_id>/duzenle", methods=["GET", "POST"])
def kulup_form(kulup_id=None):
    mevcut = _kulup_veya_404(kulup_id) if kulup_id else None
    hata = None

    if request.method == "POST":
        f = request.form
        kulup = {
            "isim": f.get("isim", "").strip(),
            "kategori": f.get("kategori", "").strip(),
            "eposta": f.get("eposta", "").strip(),
            "aciklama": f.get("aciklama", "").strip(),
            "baskan": f.get("baskan", "").strip(),
            "durum": "pasif" if f.get("durum") == "pasif" else "aktif",
            "danisman_id": f.get("danisman_id", ""),
            "yonetim_kurulu": [
                {"gorev": g.strip(), "isim": i.strip()}
                for g, i in zip(f.getlist("yk_gorev"), f.getlist("yk_isim"))
                if g.strip() and i.strip()
            ],
        }
        logo = request.files.get("logo")

        if not kulup["isim"] or not kulup["kategori"]:
            hata = "Kulüp adı ve kategori boş bırakılamaz."
        elif kulup["danisman_id"] and veri.danisman_getir(kulup["danisman_id"]) is None:
            hata = "Seçilen danışman bulunamadı."
        elif kulup["eposta"] and not EPOSTA_DESENI.fullmatch(kulup["eposta"]):
            hata = "Kulüp e-posta adresi geçersiz."
        elif logo and logo.filename and _uzanti(logo.filename) not in LOGO_UZANTILARI:
            hata = "Logo PNG, JPG veya WEBP olmalı."
        else:
            kulup["id"] = mevcut["id"] if mevcut else _kulup_id_uret(kulup["isim"])
            if logo and logo.filename:
                dosya = f"{kulup['id']}_logo_{uuid4().hex[:6]}.{_uzanti(logo.filename)}"
                logo.save(os.path.join(LOGO_KLASORU, dosya))
                kulup["logo"] = dosya
            elif not mevcut:
                kulup["logo"] = ""
            # Başkan alanı esastır: kuruldaki "Başkan" satırı ona göre güncellenir.
            kurul = [u for u in kulup["yonetim_kurulu"] if u["gorev"] != "Başkan"]
            if kulup["baskan"]:
                kurul.insert(0, {"gorev": "Başkan", "isim": kulup["baskan"]})
            kulup["yonetim_kurulu"] = kurul

            veri.kulup_kaydet(kulup)
            veri.gorev_degisikliklerini_kaydet(
                kulup["id"], veri.kulup_gorevleri(mevcut) if mevcut else set(),
                veri.kulup_gorevleri(kulup), session.get("ad_soyad", ""))
            _islem("Güncelleme" if mevcut else "Ekleme", "kulup", kulup["isim"], kulup["id"], kulup["id"],
                   eski=mevcut or {}, yeni=kulup)
            flash(f"{kulup['isim']} {'güncellendi' if mevcut else 'eklendi'}.", "basarili")
            return redirect(url_for("kulup_detay", kulup_id=kulup["id"]))

        # Hatalı gönderimde formu girilen değerlerle tekrar göster.
        mevcut = {**(mevcut or {}), **kulup}

    return render_template(
        "kulup_form.html",
        kulup=mevcut,
        hata=hata,
        kategoriler=veri.kategorileri_getir(),
        danismanlar=veri.danismanlari_getir(),
        yk_gorevleri=YK_GOREVLERI,
        # Yeni kulüpte standart görevler boş isimle hazır gelir.
        # Başkan ayrı alanda düzenlendiği için kurul listesinde gösterilmez.
        yk_satirlari=[u for u in (mevcut or {}).get("yonetim_kurulu", []) if u["gorev"] != "Başkan"]
        or [{"gorev": g, "isim": ""} for g in YK_GOREVLERI if g != "Başkan"],
    )


@app.route("/kulupler/<kulup_id>/sil", methods=["POST"])
def kulup_sil(kulup_id):
    """Kulübü silmez, arşive alır (üyelikleri, belgeleri ve başvuruları onunla birlikte gizlenir)."""
    kulup = _kulup_veya_404(kulup_id)
    veri.kulup_arsiv_ayarla(kulup_id, session.get("ad_soyad", ""))
    _islem("Arşive alma", "kulup", kulup["isim"], kulup_id, kulup_id)
    flash(f"{kulup['isim']} arşive alındı. Arşiv sayfasından geri getirilebilir.", "basarili")
    return redirect(url_for("kulupler"))


@app.route("/kulupler/<kulup_id>")
def kulup_detay(kulup_id):
    kulup = _kulup_veya_404(kulup_id)
    return render_template(
        "kulup_detay.html",
        kulup=kulup,
        uyeler=veri.uyeleri_getir(kulup_id),
        belgeler=veri.belgeleri_getir(kulup_id),
        etkinlikler=veri.basvurulari_getir(kulup_id=kulup_id),
        gorev_gecmisi=veri.gorev_gecmisi_getir(kulup_id),
    )


# ---------------------------------------------------------------------------
# Kulüp üyeleri
# ---------------------------------------------------------------------------

@app.route("/kulupler/<kulup_id>/uyeler/yeni", methods=["GET", "POST"])
@app.route("/kulupler/<kulup_id>/uyeler/<ogrenci_no>/duzenle", methods=["GET", "POST"])
def uye_form(kulup_id, ogrenci_no=None):
    kulup = _kulup_veya_404(kulup_id)
    mevcut = None
    if ogrenci_no:
        mevcut = next((u for u in veri.uyeleri_getir(kulup_id) if u["ogrenci_no"] == ogrenci_no), None)
        if mevcut is None:
            abort(404)
    hata = None

    if request.method == "POST":
        uye = {alan: request.form.get(alan, "").strip() for alan in veri.OGRENCI_ALANLARI}
        uye["ogrenci_no"] = request.form.get("ogrenci_no", "").strip()
        uye["katilim_tarihi"] = request.form.get("katilim_tarihi", "").strip()

        if not uye["ogrenci_no"] or not uye["ad_soyad"]:
            hata = "Öğrenci numarası ve ad soyad boş bırakılamaz."
        elif not re.fullmatch(r"[0-9A-Za-zçğıöşüÇĞİÖŞÜ]{1,32}", uye["ogrenci_no"]):
            hata = "Öğrenci numarası geçersiz."
        elif uye["eposta"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", uye["eposta"]):
            hata = "E-posta adresi geçersiz."
        elif uye["ogrenci_no"] != ogrenci_no and veri.uye_mi(kulup_id, uye["ogrenci_no"]):
            hata = "Bu öğrenci zaten bu kulübün üyesi."
        else:
            try:
                if uye["katilim_tarihi"]:
                    datetime.strptime(uye["katilim_tarihi"], "%Y-%m-%d")
            except ValueError:
                hata = "Katılım tarihi geçersiz."
            else:
                veri.uye_kaydet(kulup_id, uye["ogrenci_no"], uye, uye["katilim_tarihi"], ogrenci_no)
                _islem("Güncelleme" if mevcut else "Ekleme", "uye", f"{uye['ad_soyad']} ({uye['ogrenci_no']})",
                       uye["ogrenci_no"], kulup_id, eski=mevcut or {}, yeni=uye)
                flash(f"{uye['ad_soyad']} {'güncellendi' if mevcut else 'eklendi'}.", "basarili")
                return redirect(url_for("kulup_detay", kulup_id=kulup_id) + "#uyeler")
        mevcut = {**(mevcut or {}), **uye}

    return render_template(
        "uye_form.html",
        kulup=kulup,
        uye=mevcut,
        duzenleme=ogrenci_no is not None,
        hata=hata,
        bugun=date.today().isoformat(),
    )


@app.route("/kulupler/<kulup_id>/uyeler/<ogrenci_no>/sil", methods=["POST"])
def uye_sil(kulup_id, ogrenci_no):
    kulup = _kulup_veya_404(kulup_id)
    uye = next((u for u in veri.uyeleri_getir(kulup_id) if u["ogrenci_no"] == ogrenci_no), None)
    if uye is None:
        abort(404)
    veri.uye_sil(kulup_id, ogrenci_no)
    _islem("Silme", "uye", f"{uye['ad_soyad'] or ogrenci_no} ({ogrenci_no})", ogrenci_no, kulup_id,
           aciklama=f"{kulup['isim']} üyeliğinden çıkarıldı.")
    flash(f"{ogrenci_no} numaralı öğrenci kulüpten çıkarıldı.", "basarili")
    return redirect(url_for("kulup_detay", kulup_id=kulup_id) + "#uyeler")


UYE_SUTUNLARI = [
    ("ad_soyad", "Ad Soyad"), ("ogrenci_no", "Öğrenci No"), ("bolum", "Bölüm"),
    ("sinif", "Sınıf"), ("telefon", "Telefon"), ("eposta", "E-posta"),
    ("katilim_tarihi", "Katılım Tarihi"),
]


def _dosya_adi(kulup):
    return f"{kulup['id']}-uyeler-{date.today().isoformat()}"


@app.route("/kulupler/<kulup_id>/uyeler.xlsx")
def uyeler_excel(kulup_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    kulup = _kulup_veya_404(kulup_id)
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Üyeler"
    sayfa.append([baslik for _, baslik in UYE_SUTUNLARI])
    for hucre in sayfa[1]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="1B2F8F")
    for uye in veri.uyeleri_getir(kulup_id):
        sayfa.append([uye[alan] for alan, _ in UYE_SUTUNLARI])
    for sutun, genislik in zip("ABCDEFG", (26, 14, 28, 8, 16, 30, 15)):
        sayfa.column_dimensions[sutun].width = genislik
    sayfa.freeze_panes = "A2"

    cikti = io.BytesIO()
    kitap.save(cikti)
    cikti.seek(0)
    return send_file(cikti, as_attachment=True, download_name=_dosya_adi(kulup) + ".xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _pdf_fontu():
    """Türkçe karakterleri destekleyen bir TTF yazı tipi kaydeder; bulamazsa Helvetica döner."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    adaylar = [
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
    ]
    for normal, kalin in adaylar:
        if os.path.exists(normal) and os.path.exists(kalin):
            if "PanelFont" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("PanelFont", normal))
                pdfmetrics.registerFont(TTFont("PanelFont-Bold", kalin))
            return "PanelFont", "PanelFont-Bold"
    return "Helvetica", "Helvetica-Bold"


@app.route("/kulupler/<kulup_id>/uyeler.pdf")
def uyeler_pdf(kulup_id):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    kulup = _kulup_veya_404(kulup_id)
    uyeler = veri.uyeleri_getir(kulup_id)
    font, font_kalin = _pdf_fontu()

    cikti = io.BytesIO()
    belge = SimpleDocTemplate(cikti, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                              topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=f"{kulup['isim']} Üye Listesi")
    baslik = ParagraphStyle("baslik", fontName=font_kalin, fontSize=15, leading=19, textColor=colors.HexColor("#12205F"))
    alt = ParagraphStyle("alt", fontName=font, fontSize=9, leading=12, textColor=colors.HexColor("#757E8C"))
    hucre = ParagraphStyle("hucre", fontName=font, fontSize=8, leading=10)

    satirlar = [[baslik_ for _, baslik_ in UYE_SUTUNLARI]]
    for uye in uyeler:
        satirlar.append([Paragraph(str(uye[alan] or "—"), hucre) for alan, _ in UYE_SUTUNLARI])
    tablo = Table(satirlar, repeatRows=1,
                  colWidths=[5 * cm, 2.6 * cm, 5 * cm, 1.4 * cm, 3 * cm, 6 * cm, 2.6 * cm])
    tablo.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), font_kalin),
        ("FONTSIZE", (0, 0), (-1, 0), 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B2F8F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F5FA")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DCE6")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    belge.build([
        Paragraph(f"{kulup['isim']} — Üye Listesi", baslik),
        Paragraph(f"Altınbaş Üniversitesi Öğrenci Dekanlığı · {len(uyeler)} üye · "
                  f"{date.today().strftime('%d.%m.%Y')}", alt),
        Spacer(1, 0.4 * cm),
        tablo,
    ])
    cikti.seek(0)
    return send_file(cikti, as_attachment=True, download_name=_dosya_adi(kulup) + ".pdf",
                     mimetype="application/pdf")


# ---------------------------------------------------------------------------
# Tüm üyeler
# ---------------------------------------------------------------------------

@app.route("/uyeler")
def uyeler():
    return render_template(
        "uyeler.html",
        uyeler=veri.tum_uyeleri_getir(),
        kulupler=veri.kulupleri_getir(),
        secili_kulup=request.args.get("kulup", ""),
    )


# ---------------------------------------------------------------------------
# Danışmanlar
# ---------------------------------------------------------------------------

UNVANLAR = ["Prof. Dr.", "Doç. Dr.", "Dr. Öğr. Üyesi", "Öğr. Gör. Dr.", "Öğr. Gör.", "Arş. Gör. Dr.", "Arş. Gör."]


def _danisman_veya_404(danisman_id):
    danisman = veri.danisman_getir(danisman_id)
    if danisman is None:
        abort(404)
    return danisman


@app.route("/danismanlar")
def danismanlar():
    return render_template("danismanlar.html", danismanlar=veri.danismanlari_getir())


@app.route("/danismanlar/<danisman_id>")
def danisman_detay(danisman_id):
    danisman = _danisman_veya_404(danisman_id)
    kulupler = [k for k in veri.kulupleri_getir() if k["danisman_id"] == danisman_id]
    return render_template("danisman_detay.html", danisman=danisman, kulupler=kulupler)


@app.route("/danismanlar/yeni", methods=["GET", "POST"])
@app.route("/danismanlar/<danisman_id>/duzenle", methods=["GET", "POST"])
def danisman_form(danisman_id=None):
    mevcut = _danisman_veya_404(danisman_id) if danisman_id else None
    hata = None

    if request.method == "POST":
        danisman = {alan: request.form.get(alan, "").strip() for alan in veri.DANISMAN_ALANLARI}
        if not danisman["ad_soyad"]:
            hata = "Ad soyad boş bırakılamaz."
        elif danisman["eposta"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", danisman["eposta"]):
            hata = "E-posta adresi geçersiz."
        else:
            danisman["id"] = mevcut["id"] if mevcut else veri.yeni_id()
            veri.danisman_kaydet(danisman)
            _islem("Güncelleme" if mevcut else "Ekleme", "danisman", veri.danisman_tam_adi(danisman),
                   danisman["id"], eski=mevcut or {}, yeni=danisman)
            flash(f"{veri.danisman_tam_adi(danisman)} {'güncellendi' if mevcut else 'eklendi'}.", "basarili")
            # Kulüp formundan "yeni danışman" ile gelindiyse oraya geri dön.
            geri = request.args.get("geri")
            if geri:
                return redirect(_guvenli_yonlendirme(geri))
            return redirect(url_for("danisman_detay", danisman_id=danisman["id"]))
        mevcut = {**(mevcut or {}), **danisman}

    return render_template("danisman_form.html", danisman=mevcut, duzenleme=danisman_id is not None,
                           hata=hata, unvanlar=UNVANLAR)


@app.route("/danismanlar/<danisman_id>/sil", methods=["POST"])
def danisman_sil(danisman_id):
    danisman = _danisman_veya_404(danisman_id)
    veri.danisman_sil(danisman_id)
    _islem("Silme", "danisman", danisman["tam_ad"], danisman_id,
           aciklama=("Danışmanlık yaptığı kulüpler: " + ", ".join(k["isim"] for k in danisman["kulupler"]))
           if danisman["kulupler"] else "")
    flash(f"{danisman['tam_ad']} silindi."
          + (f" {len(danisman['kulupler'])} kulübün danışman alanı boşaltıldı." if danisman["kulupler"] else ""),
          "basarili")
    return redirect(url_for("danismanlar"))


# ---------------------------------------------------------------------------
# Dönemler
# ---------------------------------------------------------------------------

def _donem_formu():
    """Formdan dönem bilgisini okur; hata varsa (None, mesaj) döner."""
    donem = {
        "ad": request.form.get("ad", "").strip(),
        "baslangic": request.form.get("baslangic", "").strip(),
        "bitis": request.form.get("bitis", "").strip(),
    }
    if not all(donem.values()):
        return None, "Dönem adı, başlangıç ve bitiş tarihi zorunludur."
    try:
        if datetime.strptime(donem["bitis"], "%Y-%m-%d") <= datetime.strptime(donem["baslangic"], "%Y-%m-%d"):
            return None, "Bitiş tarihi başlangıçtan sonra olmalı."
    except ValueError:
        return None, "Tarihler geçersiz."
    return donem, None


@app.route("/donemler", methods=["GET", "POST"])
def donemler():
    hata = None
    if request.method == "POST":
        donem, hata = _donem_formu()
        if donem:
            if any(d["ad"].lower() == donem["ad"].lower() for d in veri.donemleri_getir()):
                hata = "Bu adda bir dönem zaten var."
            else:
                donem["id"] = veri.yeni_id()
                veri.donem_kaydet(donem)
                if request.form.get("aktif_yap"):
                    veri.donem_aktif_yap(donem["id"])
                _islem("Ekleme", "donem", donem["ad"], donem["id"], yeni=donem,
                       aciklama="Aktif dönem yapıldı." if request.form.get("aktif_yap") else "")
                flash(f"{donem['ad']} eklendi.", "basarili")
                return redirect(url_for("donemler"))

    donem_listesi = veri.donemleri_getir()
    for donem in donem_listesi:
        donem["kayit_sayisi"] = veri.donem_kullanim_sayisi(donem["id"])
    return render_template("donemler.html", donemler=donem_listesi, hata=hata, form=request.form)


@app.route("/donemler/<donem_id>/duzenle", methods=["GET", "POST"])
def donem_duzenle(donem_id):
    mevcut = veri.donem_getir(donem_id)
    if mevcut is None:
        abort(404)
    hata = None
    if request.method == "POST":
        donem, hata = _donem_formu()
        if donem:
            if any(d["ad"].lower() == donem["ad"].lower() and d["id"] != donem_id for d in veri.donemleri_getir()):
                hata = "Bu adda bir dönem zaten var."
            else:
                veri.donem_kaydet({**donem, "id": donem_id})
                _islem("Güncelleme", "donem", donem["ad"], donem_id, eski=mevcut, yeni=donem)
                flash(f"{donem['ad']} güncellendi.", "basarili")
                return redirect(url_for("donemler"))
        mevcut = {**mevcut, **request.form.to_dict()}
    return render_template("donem_form.html", donem=mevcut, hata=hata)


@app.route("/donemler/<donem_id>/aktif", methods=["POST"])
def donem_aktif_yap(donem_id):
    donem = veri.donem_getir(donem_id)
    if donem is None:
        abort(404)
    onceki = veri.aktif_donem()
    veri.donem_aktif_yap(donem_id)
    _islem("Güncelleme", "donem", donem["ad"], donem_id,
           aciklama=f"Aktif dönem yapıldı (önceki: {onceki['ad'] if onceki else '—'}).")
    flash(f"Aktif dönem: {donem['ad']}.", "basarili")
    return redirect(url_for("donemler"))


@app.route("/donemler/<donem_id>/sil", methods=["POST"])
def donem_sil(donem_id):
    donem = veri.donem_getir(donem_id)
    if donem is None:
        abort(404)
    if donem["aktif"]:
        flash("Aktif dönem silinemez. Önce başka bir dönemi aktif yapın.", "hata")
    elif veri.donem_kullanim_sayisi(donem_id):
        flash("Bu döneme bağlı belge veya zorunlu evrak olduğu için dönem silinemez.", "hata")
    else:
        veri.donem_sil(donem_id)
        _islem("Silme", "donem", donem["ad"], donem_id)
        flash(f"{donem['ad']} silindi.", "basarili")
    return redirect(url_for("donemler"))


# ---------------------------------------------------------------------------
# Etkinlik başvuruları
# ---------------------------------------------------------------------------

BASVURU_KAYNAKLARI = {"e_posta": "E-posta ile geldi", "elden": "Elden teslim", "panel": "Panelden girildi",
                      "baskan": "Kulüp başkanı panelden girdi", "eski_site": "Eski siteden aktarıldı"}


def _basvuru_dosyasini_sil(dosya):
    yol = os.path.join(BASVURU_KLASORU, dosya["dosya"])
    if os.path.exists(yol):
        os.remove(yol)


def _basvuru_veya_404(basvuru_id):
    basvuru = veri.basvuru_getir(basvuru_id)
    if basvuru is None:
        abort(404)
    return basvuru


def _basvuru_formu(kulup_isimleri):
    """Formdan başvuru alanlarını okur; (başvuru, hata) döndürür."""
    f = request.form
    b = {
        "kulup_id": f.get("kulup_id", ""),
        "ad": f.get("ad", "").strip(),
        "aciklama": f.get("aciklama", "").strip(),
        "tarih_saat": f.get("tarih_saat", "").strip(),
        "yer": f.get("yer", "").strip(),
        "konusmaci_var": f.get("konusmaci_var") == "evet",
        "konusmaci": f.get("konusmaci", "").strip(),
        "butce_aciklama": f.get("butce_aciklama", "").strip(),
        "kaynak": f.get("kaynak") if f.get("kaynak") in ("e_posta", "elden", "panel") else "panel",
        "saat_belirsiz": False,
    }
    if not b["konusmaci_var"]:
        b["konusmaci"] = ""
    try:
        b["katilimci"] = int(f.get("katilimci") or 0)
        b["butce"] = float((f.get("butce") or "0").replace(",", "."))
    except ValueError:
        return b, "Katılımcı sayısı ve bütçe sayı olmalı."

    if b["kulup_id"] not in kulup_isimleri:
        return b, "Lütfen bir kulüp seçin."
    if not b["ad"] or not b["tarih_saat"] or not b["yer"]:
        return b, "Etkinlik adı, tarih-saat ve yer zorunludur."
    try:
        datetime.strptime(b["tarih_saat"], "%Y-%m-%dT%H:%M")
    except ValueError:
        return b, "Tarih-saat geçersiz."
    if b["katilimci"] < 0 or b["butce"] < 0:
        return b, "Katılımcı sayısı ve bütçe negatif olamaz."
    if b["konusmaci_var"] and not b["konusmaci"]:
        return b, "Konuşmacı/misafir bilgisini yazın."
    return b, None


def _basvuru_dosyalarini_kaydet():
    """Yüklenen ek dosyaları kaydeder; (dosya listesi, hata) döndürür."""
    dosyalar = [d for d in request.files.getlist("ek_dosya") if d and d.filename]
    if any(_uzanti(d.filename) not in BELGE_UZANTILARI for d in dosyalar):
        return [], "Ek dosya türü desteklenmiyor."
    kayitlar = []
    os.makedirs(BASVURU_KLASORU, exist_ok=True)
    for dosya in dosyalar:
        dosya_id = uuid4().hex
        kayit_adi = f"{dosya_id}.{_uzanti(dosya.filename)}"
        dosya.save(os.path.join(BASVURU_KLASORU, kayit_adi))
        kayitlar.append({
            "id": dosya_id,
            "dosya": kayit_adi,
            "orijinal_ad": secure_filename(dosya.filename) or kayit_adi,
            "boyut": os.path.getsize(os.path.join(BASVURU_KLASORU, kayit_adi)),
        })
    return kayitlar, None


def _kullanicinin_kulupleri():
    """Kulüp başkanı için sadece kendi kulübü, diğer roller için tüm kulüpler."""
    kulupler = veri.kulupleri_getir()
    if _baskan_kulubu():
        kulupler = [k for k in kulupler if k["id"] == _baskan_kulubu()]
    return kulupler


@app.route("/etkinlik-basvurulari")
def etkinlik_basvurulari():
    basvurular = veri.basvurulari_getir(kulup_id=_baskan_kulubu())
    sayilar = {d: sum(1 for b in basvurular if b["durum"] == d) for d in veri.BASVURU_DURUMLARI}
    return render_template(
        "basvurular.html",
        basvurular=basvurular,
        sayilar=sayilar,
        durumlar=veri.BASVURU_DURUMLARI,
        kulupler=_kullanicinin_kulupleri(),
        donemler=veri.donemleri_getir(),
        secili_durum=request.args.get("durum", ""),
        secili_kulup=request.args.get("kulup", ""),
    )


@app.route("/etkinlik-basvurulari/yeni", methods=["GET", "POST"])
@app.route("/etkinlik-basvurulari/<basvuru_id>/duzenle", methods=["GET", "POST"])
def basvuru_form(basvuru_id=None):
    mevcut = _basvuru_veya_404(basvuru_id) if basvuru_id else None
    # Kulüp başkanı karar verilmiş başvuruyu değiştiremez; sadece bekleyen/revizyondakini düzenler.
    if mevcut and _baskan_kulubu() and mevcut["durum"] not in ("beklemede", "revizyon"):
        abort(403)
    kulupler = _kullanicinin_kulupleri()
    kulup_isimleri = {k["id"]: k["isim"] for k in kulupler}
    hata = None

    if request.method == "POST":
        basvuru, hata = _basvuru_formu(kulup_isimleri)
        if _baskan_kulubu():
            basvuru["kaynak"] = "baskan"
        if not hata:
            yeni_dosyalar, hata = _basvuru_dosyalarini_kaydet()
        if not hata:
            kullanici = session.get("ad_soyad", "")
            if mevcut:
                basvuru["id"] = mevcut["id"]
                basvuru["kaynak"] = mevcut.get("kaynak", basvuru["kaynak"])
                basvuru["ek_dosyalar"] = mevcut.get("ek_dosyalar", []) + yeni_dosyalar
                veri.basvuru_kaydet(basvuru)
                # Revizyon istenen başvuru düzeltilince yeniden değerlendirmeye düşer.
                if mevcut["durum"] == "revizyon":
                    veri.basvuru_gecmise_ekle(mevcut["id"], kullanici, "beklemede", "Başvuru revize edildi.")
                _islem("Güncelleme", "basvuru", basvuru["ad"], basvuru["id"], basvuru["kulup_id"],
                       eski=mevcut, yeni=basvuru,
                       aciklama=f"{len(yeni_dosyalar)} ek dosya eklendi." if yeni_dosyalar else "")
                flash("Başvuru güncellendi.", "basarili")
            else:
                basvuru.update({
                    "id": veri.yeni_id(),
                    "ek_dosyalar": yeni_dosyalar,
                    "giren": kullanici,
                    "olusturma": datetime.now().strftime("%Y-%m-%d %H:%M"),
                    "durum": "beklemede",
                    "gecmis": [],
                })
                veri.basvuru_kaydet(basvuru)
                veri.basvuru_gecmise_ekle(basvuru["id"], kullanici, "beklemede",
                                          f"Başvuru kaydedildi ({BASVURU_KAYNAKLARI[basvuru['kaynak']].lower()}).")
                _islem("Ekleme", "basvuru", basvuru["ad"], basvuru["id"], basvuru["kulup_id"], yeni=basvuru)
                flash("Başvuru kaydedildi, onay bekliyor.", "basarili")
            return redirect(url_for("basvuru_detay", basvuru_id=basvuru["id"]))
        mevcut = {**(mevcut or {}), **basvuru}
    elif not mevcut:
        mevcut = {"kulup_id": _baskan_kulubu() or request.args.get("kulup", ""), "kaynak": "e_posta"}

    return render_template(
        "basvuru_form.html",
        basvuru=mevcut,
        duzenleme=basvuru_id is not None,
        kulupler=kulupler,
        kaynaklar={k: v for k, v in BASVURU_KAYNAKLARI.items() if k not in ("eski_site", "baskan")},
        hata=hata,
        uzantilar=", ".join(sorted(BELGE_UZANTILARI)),
    )


@app.route("/etkinlik-basvurulari/<basvuru_id>")
def basvuru_detay(basvuru_id):
    basvuru = _basvuru_veya_404(basvuru_id)
    return render_template("basvuru_detay.html", basvuru=basvuru, durumlar=veri.BASVURU_DURUMLARI,
                           kaynaklar=BASVURU_KAYNAKLARI)


@app.route("/etkinlik-basvurulari/<basvuru_id>/karar", methods=["POST"])
def basvuru_karar(basvuru_id):
    basvuru = _basvuru_veya_404(basvuru_id)
    durum = request.form.get("durum", "")
    aciklama = request.form.get("aciklama", "").strip()
    if durum not in ("onaylandi", "reddedildi", "revizyon", "beklemede"):
        flash("Geçersiz karar.", "hata")
    elif durum in ("reddedildi", "revizyon") and not aciklama:
        flash("Red ve revizyon kararlarında açıklama yazmak zorunludur.", "hata")
    elif durum == basvuru["durum"] and not aciklama:
        flash("Başvuru zaten bu durumda.", "hata")
    else:
        veri.basvuru_gecmise_ekle(basvuru_id, session.get("ad_soyad", ""), durum, aciklama)
        _islem("Karar", "basvuru", basvuru["ad"], basvuru_id, basvuru["kulup_id"],
               eski={"durum": basvuru["durum_ad"]}, yeni={"durum": veri.BASVURU_DURUMLARI[durum]}, aciklama=aciklama)
        flash(f"Başvuru durumu: {veri.BASVURU_DURUMLARI[durum]}.", "basarili")
    return redirect(url_for("basvuru_detay", basvuru_id=basvuru_id))


@app.route("/etkinlik-basvurulari/<basvuru_id>/dosya/<dosya_id>")
def basvuru_dosya(basvuru_id, dosya_id):
    basvuru = _basvuru_veya_404(basvuru_id)
    dosya = next((d for d in basvuru.get("ek_dosyalar", []) if d["id"] == dosya_id), None)
    if dosya is None:
        abort(404)
    return send_from_directory(BASVURU_KLASORU, dosya["dosya"], as_attachment=True,
                               download_name=dosya["orijinal_ad"])


@app.route("/etkinlik-basvurulari/<basvuru_id>/sil", methods=["POST"])
def basvuru_sil(basvuru_id):
    """Başvuruyu silmez, arşive alır."""
    basvuru = _basvuru_veya_404(basvuru_id)
    veri.basvuru_arsiv_ayarla(basvuru_id, session.get("ad_soyad", ""))
    _islem("Arşive alma", "basvuru", basvuru["ad"], basvuru_id, basvuru["kulup_id"])
    flash(f"{basvuru['ad']} başvurusu arşive alındı.", "basarili")
    return redirect(url_for("etkinlik_basvurulari"))


# ---------------------------------------------------------------------------
# Zorunlu evraklar ve evrak takibi
# ---------------------------------------------------------------------------

def _donem_veya_404(donem_id):
    donem = veri.donem_getir(donem_id)
    if donem is None:
        abort(404)
    return donem


def _secili_donem():
    """?donem= parametresindeki dönem; yoksa aktif dönem (o da yoksa None)."""
    return veri.donem_getir(request.args.get("donem", "")) or veri.aktif_donem()


def _zorunlu_evrak_formu():
    evrak = {
        "tur": request.form.get("tur", ""),
        "teslim_tarihi": request.form.get("teslim_tarihi", "").strip(),
        "aciklama": request.form.get("aciklama", "").strip(),
    }
    if evrak["tur"] not in BELGE_TURLERI:
        return evrak, "Lütfen belge türünü seçin."
    try:
        date.fromisoformat(evrak["teslim_tarihi"])
    except ValueError:
        return evrak, "Geçerli bir teslim tarihi girin."
    return evrak, None


@app.route("/donemler/<donem_id>/zorunlu-evraklar", methods=["GET", "POST"])
def zorunlu_evraklar(donem_id):
    donem = _donem_veya_404(donem_id)
    hata = None
    if request.method == "POST":
        evrak, hata = _zorunlu_evrak_formu()
        if not hata and any(e["tur"] == evrak["tur"] for e in veri.zorunlu_evraklari_getir(donem_id)):
            hata = f"{evrak['tur']} bu dönem için zaten zorunlu evrak olarak tanımlı."
        if not hata:
            evrak_id = veri.yeni_id()
            veri.zorunlu_evrak_kaydet({**evrak, "id": evrak_id, "donem_id": donem_id, "elden": {}})
            _islem("Ekleme", "zorunlu_evrak", f"{evrak['tur']} ({donem['ad']})", evrak_id, yeni=evrak)
            flash(f"{evrak['tur']} zorunlu evraklara eklendi.", "basarili")
            return redirect(url_for("zorunlu_evraklar", donem_id=donem_id))
    return render_template("zorunlu_evraklar.html", donem=donem, evraklar=veri.zorunlu_evraklari_getir(donem_id),
                           belge_turleri=BELGE_TURLERI, hata=hata, form=request.form)


@app.route("/zorunlu-evraklar/<evrak_id>/duzenle", methods=["GET", "POST"])
def zorunlu_evrak_duzenle(evrak_id):
    evrak = veri.zorunlu_evrak_getir(evrak_id)
    if evrak is None:
        abort(404)
    hata = None
    if request.method == "POST":
        yeni, hata = _zorunlu_evrak_formu()
        if not hata and any(e["tur"] == yeni["tur"] and e["id"] != evrak_id
                            for e in veri.zorunlu_evraklari_getir(evrak["donem_id"])):
            hata = f"{yeni['tur']} bu dönem için zaten zorunlu evrak olarak tanımlı."
        if not hata:
            veri.zorunlu_evrak_kaydet({**yeni, "id": evrak_id})
            _islem("Güncelleme", "zorunlu_evrak", yeni["tur"], evrak_id, eski=evrak, yeni=yeni)
            flash("Zorunlu evrak güncellendi.", "basarili")
            return redirect(url_for("zorunlu_evraklar", donem_id=evrak["donem_id"]))
        evrak = {**evrak, **yeni}
    return render_template("zorunlu_evrak_form.html", evrak=evrak, donem=veri.donem_getir(evrak["donem_id"]),
                           belge_turleri=BELGE_TURLERI, hata=hata)


@app.route("/zorunlu-evraklar/<evrak_id>/sil", methods=["POST"])
def zorunlu_evrak_sil(evrak_id):
    evrak = veri.zorunlu_evrak_getir(evrak_id)
    if evrak is None:
        abort(404)
    veri.zorunlu_evrak_sil(evrak_id)
    donem = veri.donem_getir(evrak["donem_id"])
    _islem("Silme", "zorunlu_evrak", f"{evrak['tur']} ({donem['ad'] if donem else ''})", evrak_id)
    flash(f"{evrak['tur']} zorunlu evraklardan çıkarıldı. Yüklenmiş belgeler silinmedi.", "basarili")
    return redirect(url_for("zorunlu_evraklar", donem_id=evrak["donem_id"]))


@app.route("/evrak-takibi")
def evrak_takibi():
    donem = _secili_donem()
    pasifler_dahil = request.args.get("pasif") == "1"
    evraklar, satirlar = [], []
    if donem:
        kulupler = [k for k in veri.kulupleri_getir() if pasifler_dahil or k["durum"] == "aktif"]
        evraklar, satirlar = veri.evrak_durumlari(donem["id"], kulupler=kulupler)
    toplam = {d: sum(s["sayilar"][d] for s in satirlar) for d in veri.EVRAK_DURUMLARI}
    return render_template("evrak_takibi.html", donem=donem, donemler=veri.donemleri_getir(), evraklar=evraklar,
                           satirlar=satirlar, toplam=toplam, durumlar=veri.EVRAK_DURUMLARI,
                           pasifler_dahil=pasifler_dahil, yaklasiyor_gun=veri.YAKLASIYOR_GUN)


@app.route("/evrak-takibi/elden", methods=["POST"])
def elden_teslim():
    evrak = veri.zorunlu_evrak_getir(request.form.get("evrak_id", ""))
    kulup = veri.kulup_getir(request.form.get("kulup_id", ""))
    if evrak is None or kulup is None:
        abort(404)
    if veri.elden_teslim_ayarla(evrak["id"], kulup["id"], session.get("ad_soyad", "")):
        mesaj = "elden teslim alındı olarak işaretlendi"
    else:
        mesaj = "elden teslim işareti kaldırıldı"
    _islem("Güncelleme", "zorunlu_evrak", evrak["tur"], evrak["id"], kulup["id"], aciklama=f"{kulup['isim']}: {mesaj}.")
    flash(f"{kulup['isim']} · {evrak['tur']}: {mesaj}.", "basarili")
    return redirect(url_for("evrak_takibi", donem=evrak["donem_id"]) + f"#kulup-{kulup['id']}")


# ---------------------------------------------------------------------------
# Kulüp karnesi
# ---------------------------------------------------------------------------

@app.route("/karne")
def karne():
    donem = _secili_donem()
    pasifler_dahil = request.args.get("pasif") == "1"
    karneler = []
    if donem:
        kulupler = [k for k in veri.kulupleri_getir() if pasifler_dahil or k["durum"] == "aktif"]
        karneler = veri.donem_karneleri(donem["id"], kulupler)
    puanli = [k["puan"] for k in karneler if k["puan"] is not None]
    return render_template("karne.html", donem=donem, donemler=veri.donemleri_getir(), karneler=karneler,
                           pasifler_dahil=pasifler_dahil, degerlendirilen=len(puanli),
                           ortalama=round(sum(puanli) / len(puanli), 2) if puanli else None)


@app.route("/kulupler/<kulup_id>/karne", methods=["GET", "POST"])
def kulup_karne(kulup_id):
    kulup = _kulup_veya_404(kulup_id)
    donem = _secili_donem()
    if donem is None:
        flash("Karne için önce bir dönem tanımlayın.", "hata")
        return redirect(url_for("donemler"))
    hata = None
    if request.method == "POST":
        puan = request.form.get("puan", "")
        notu = request.form.get("not", "").strip()
        if puan not in {"1", "2", "3", "4", "5"}:
            hata = "1 ile 5 arasında bir puan seçin."
        else:
            eski = veri.karne_puanlari(donem["id"]).get(kulup_id) or {}
            veri.karne_kaydet(donem["id"], kulup_id, int(puan), notu, session.get("ad_soyad", ""))
            _islem("Güncelleme" if eski else "Ekleme", "karne", f"{kulup['isim']} · {donem['ad']}", kulup_id, kulup_id,
                   eski=eski, yeni={"puan": int(puan), "not": notu})
            flash(f"{kulup['isim']} · {donem['ad']} karnesi kaydedildi.", "basarili")
            return redirect(url_for("karne", donem=donem["id"]))
    karne_bilgisi = next(k for k in veri.donem_karneleri(donem["id"], [kulup]))
    return render_template("kulup_karne.html", kulup=kulup, donem=donem, donemler=veri.donemleri_getir(),
                           karne=karne_bilgisi, hata=hata)


# ---------------------------------------------------------------------------
# Raporlar
# ---------------------------------------------------------------------------

RAPOR_SUTUNLARI = [
    ("kulup", "Kulüp"), ("kategori", "Kategori"), ("durum", "Durum"), ("baskan", "Başkan"),
    ("danisman", "Danışman"), ("uye", "Üye"), ("etkinlik", "Etkinlik"), ("evrak", "Evrak"), ("puan", "Puan"),
]


def _rapor_satirlari(donem):
    """Dönem raporunun tüm kulüpler için satırları (isim sırasıyla)."""
    satirlar = []
    for k in sorted(veri.donem_karneleri(donem["id"]), key=lambda k: k["kulup"]["isim"].lower()):
        kulup = k["kulup"]
        satirlar.append({
            "kulup": kulup["isim"],
            "kategori": kulup["kategori"],
            "durum": "Aktif" if kulup["durum"] == "aktif" else "Pasif",
            "baskan": kulup["baskan"] or "Atanmamış",
            "danisman": kulup["danisman"]["isim"] or "—",
            "uye": k["uye"],
            "etkinlik": k["etkinlik"],
            "evrak": f"{k['evrak_teslim']}/{k['evrak_toplam']}" if k["evrak_toplam"] else "—",
            "puan": k["puan"] if k["puan"] is not None else "—",
        })
    return satirlar


def _grafik_verileri(donem):
    kulupler = veri.kulupleri_getir()
    kategoriler = {}
    for k in kulupler:
        kategoriler[k["kategori"]] = kategoriler.get(k["kategori"], 0) + 1

    # Dönemin aylarına göre onaylı etkinlik sayısı
    aylar = []
    ay = date.fromisoformat(donem["baslangic"]).replace(day=1)
    son = date.fromisoformat(donem["bitis"])
    while ay <= son:
        aylar.append(ay)
        ay = (ay.replace(day=28) + timedelta(days=4)).replace(day=1)
    sayac = {a.strftime("%Y-%m"): 0 for a in aylar}
    for b in veri.basvurulari_getir(durum="onaylandi"):
        if b["donem_id"] == donem["id"] and b["tarih_saat"][:7] in sayac:
            sayac[b["tarih_saat"][:7]] += 1

    aktif_kulupler = [k for k in kulupler if k["durum"] == "aktif"]
    evraklar, satirlar = veri.evrak_durumlari(donem["id"], kulupler=aktif_kulupler)
    evrak_oranlari = [
        {"ad": e["tur"], "teslim": sum(1 for s in satirlar if s["hucreler"][i]["durum"] == "teslim"),
         "toplam": len(satirlar)}
        for i, e in enumerate(evraklar)
    ]
    toplam_teslim = sum(e["teslim"] for e in evrak_oranlari)
    toplam_beklenen = sum(e["toplam"] for e in evrak_oranlari)
    return {
        "kategori": sorted(kategoriler.items(), key=lambda x: (-x[1], x[0])),
        "aktif": len(aktif_kulupler),
        "pasif": len(kulupler) - len(aktif_kulupler),
        "aylik": [(f"{AY_ADLARI[a.month - 1][:3]} {a.year}", sayac[a.strftime("%Y-%m")]) for a in aylar],
        "evrak": [(e["ad"], round(100 * e["teslim"] / e["toplam"]) if e["toplam"] else 0, e["teslim"], e["toplam"])
                  for e in evrak_oranlari],
        "evrak_genel": round(100 * toplam_teslim / toplam_beklenen) if toplam_beklenen else None,
    }


@app.route("/raporlar")
def raporlar():
    donem = _secili_donem()
    return render_template("raporlar.html", donem=donem, donemler=veri.donemleri_getir(),
                           grafik=_grafik_verileri(donem) if donem else None,
                           kulupler=veri.kulupleri_getir())


def _rapor_donemi_veya_404():
    donem = _secili_donem()
    if donem is None:
        abort(404)
    return donem


@app.route("/raporlar/donem.xlsx")
def donem_raporu_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    donem = _rapor_donemi_veya_404()
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Dönem Raporu"
    sayfa.append([f"{donem['ad']} Dönem Raporu · Altınbaş Üniversitesi Öğrenci Dekanlığı"])
    sayfa["A1"].font = Font(bold=True, size=13, color="12205F")
    sayfa.append([f"Oluşturulma: {datetime.now():%d.%m.%Y %H:%M}"])
    sayfa.append([])
    sayfa.append([b for _, b in RAPOR_SUTUNLARI])
    for hucre in sayfa[4]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="1B2F8F")
    for satir in _rapor_satirlari(donem):
        sayfa.append([satir[a] for a, _ in RAPOR_SUTUNLARI])
    for sutun, genislik in zip("ABCDEFGHI", (34, 22, 9, 22, 26, 7, 9, 9, 7)):
        sayfa.column_dimensions[sutun].width = genislik
    sayfa.freeze_panes = "A5"
    cikti = io.BytesIO()
    kitap.save(cikti)
    cikti.seek(0)
    return send_file(cikti, as_attachment=True, download_name=f"donem-raporu-{_dosya_parcasi(donem['ad'])}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _dosya_parcasi(metin):
    return re.sub(r"[^a-z0-9]+", "-", metin.translate(_TR_HARF).lower()).strip("-")


def _pdf_araclari():
    """PDF raporları için ortak yazı tipi, stiller ve tablo biçimi."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import TableStyle

    font, kalin = _pdf_fontu()
    stiller = {
        "baslik": ParagraphStyle("baslik", fontName=kalin, fontSize=15, leading=19, textColor=colors.HexColor("#12205F")),
        "alt": ParagraphStyle("alt", fontName=font, fontSize=9, leading=12, textColor=colors.HexColor("#667085")),
        "bolum": ParagraphStyle("bolum", fontName=kalin, fontSize=11, leading=14, spaceBefore=10, spaceAfter=4,
                                textColor=colors.HexColor("#12205F")),
        "metin": ParagraphStyle("metin", fontName=font, fontSize=9, leading=12),
        "hucre": ParagraphStyle("hucre", fontName=font, fontSize=8, leading=10),
    }
    tablo_stili = TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), kalin),
        ("FONTNAME", (0, 1), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B2F8F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F5FA")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DCE6")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ])
    return stiller, tablo_stili


def _pdf_yaniti(akis, dosya_adi, baslik, yatay):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate

    cikti = io.BytesIO()
    SimpleDocTemplate(cikti, pagesize=landscape(A4) if yatay else A4, title=baslik,
                      leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm).build(akis)
    cikti.seek(0)
    return send_file(cikti, as_attachment=True, download_name=dosya_adi, mimetype="application/pdf")


def _pdf_tablo(basliklar, satirlar, genislikler, stiller, tablo_stili):
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Table

    from xml.sax.saxutils import escape

    veri_satirlari = [basliklar] + [[Paragraph(escape(str(h)), stiller["hucre"]) for h in s] for s in satirlar]
    tablo = Table(veri_satirlari, repeatRows=1, colWidths=[g * cm for g in genislikler])
    tablo.setStyle(tablo_stili)
    return tablo


@app.route("/raporlar/donem.pdf")
def donem_raporu_pdf():
    from reportlab.platypus import Paragraph, Spacer
    from reportlab.lib.units import cm

    donem = _rapor_donemi_veya_404()
    stiller, tablo_stili = _pdf_araclari()
    satirlar = _rapor_satirlari(donem)
    aktif = sum(1 for s in satirlar if s["durum"] == "Aktif")
    akis = [
        Paragraph(f"{donem['ad']} Dönem Raporu", stiller["baslik"]),
        Paragraph(f"Altınbaş Üniversitesi Öğrenci Dekanlığı · {len(satirlar)} kulüp ({aktif} aktif, "
                  f"{len(satirlar) - aktif} pasif) · {datetime.now():%d.%m.%Y}", stiller["alt"]),
        Spacer(1, 0.4 * cm),
        _pdf_tablo([b for _, b in RAPOR_SUTUNLARI], [[s[a] for a, _ in RAPOR_SUTUNLARI] for s in satirlar],
                   (6, 4, 1.6, 4, 4.6, 1.3, 1.6, 1.5, 1.3), stiller, tablo_stili),
    ]
    return _pdf_yaniti(akis, f"donem-raporu-{_dosya_parcasi(donem['ad'])}.pdf", f"{donem['ad']} Dönem Raporu", True)


@app.route("/kulupler/<kulup_id>/rapor.pdf")
def kulup_raporu_pdf(kulup_id):
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Spacer

    kulup = _kulup_veya_404(kulup_id)
    donem = _rapor_donemi_veya_404()
    stiller, tablo_stili = _pdf_araclari()
    karne_bilgisi = veri.donem_karneleri(donem["id"], [kulup])[0]
    p = lambda metin, stil="metin": Paragraph(metin, stiller[stil])  # noqa: E731
    from xml.sax.saxutils import escape as e

    akis = [
        p(e(kulup["isim"]), "baslik"),
        p(f"{e(donem['ad'])} kulüp raporu · {e(kulup['kategori'])} · "
          f"{'Aktif' if kulup['durum'] == 'aktif' else 'Pasif'} · {datetime.now():%d.%m.%Y}", "alt"),
        Spacer(1, 0.3 * cm),
        p("Özet", "bolum"),
        _pdf_tablo(["Başkan", "Danışman", "Üye", "Etkinlik", "Evrak", "Puan"], [[
            kulup["baskan"] or "Atanmamış", kulup["danisman"]["isim"] or "—", karne_bilgisi["uye"],
            karne_bilgisi["etkinlik"],
            f"{karne_bilgisi['evrak_teslim']}/{karne_bilgisi['evrak_toplam']}" if karne_bilgisi["evrak_toplam"] else "—",
            f"{karne_bilgisi['puan']}/5" if karne_bilgisi["puan"] else "—",
        ]], (4.2, 4.6, 1.6, 1.8, 1.8, 1.6), stiller, tablo_stili),
    ]
    if karne_bilgisi["not"]:
        akis += [p("Dekanlık Notu", "bolum"), p(e(karne_bilgisi["not"]))]

    kurul = kulup["yonetim_kurulu"]
    akis.append(p("Yönetim Kurulu", "bolum"))
    akis.append(_pdf_tablo(["Görev", "Ad Soyad"], [[u["gorev"], u["isim"]] for u in kurul], (6, 9.6),
                           stiller, tablo_stili) if kurul else p("Kayıt yok."))

    etkinlikler = [b for b in veri.basvurulari_getir(kulup_id=kulup_id) if b["donem_id"] == donem["id"]]
    akis.append(p("Dönem Etkinlik Başvuruları", "bolum"))
    akis.append(_pdf_tablo(["Etkinlik", "Tarih", "Yer", "Durum"], [
        [b["ad"], b["tarih_saat"][:10], b["yer"] or "—", b["durum_ad"]] for b in etkinlikler
    ], (7, 2.4, 4, 3.2), stiller, tablo_stili) if etkinlikler else p("Bu dönem etkinlik başvurusu yok."))

    evraklar, satirlar = veri.evrak_durumlari(donem["id"], kulupler=[kulup])
    akis.append(p("Zorunlu Evraklar", "bolum"))
    akis.append(_pdf_tablo(["Evrak", "Son Teslim", "Durum"], [
        [h["evrak"]["tur"], h["evrak"]["teslim_tarihi"], veri.EVRAK_DURUMLARI[h["durum"]]] for h in satirlar[0]["hucreler"]
    ], (7, 3, 6.6), stiller, tablo_stili) if evraklar else p("Bu dönem için zorunlu evrak tanımlı değil."))

    gecmis = veri.gorev_gecmisi_getir(kulup_id)
    akis.append(p("Görev Geçmişi", "bolum"))
    akis.append(_pdf_tablo(["Görev", "Ad Soyad", "Başlangıç", "Bitiş"], [
        [g["gorev"], g["isim"], g["baslangic"] or "Kayıt öncesi", g["bitis"] or "Görevde"] for g in gecmis
    ], (4, 6, 3.3, 3.3), stiller, tablo_stili) if gecmis else p("Kayıt yok."))

    return _pdf_yaniti(akis, f"{kulup_id}-{_dosya_parcasi(donem['ad'])}-rapor.pdf", f"{kulup['isim']} Raporu", False)


# ---------------------------------------------------------------------------
# Takvim
# ---------------------------------------------------------------------------

AY_ADLARI = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim",
             "Kasım", "Aralık"]


@app.route("/takvim")
def takvim():
    import calendar

    try:
        ay = datetime.strptime(request.args.get("ay", ""), "%Y-%m").date()
    except ValueError:
        ay = date.today().replace(day=1)
    onceki = (ay.replace(day=1) - timedelta(days=1)).replace(day=1)
    sonraki = (ay.replace(day=28) + timedelta(days=4)).replace(day=1)

    # Gün -> o günün kayıtları
    kayitlar = {}
    for b in veri.basvurulari_getir(durum="onaylandi"):
        gun = b["tarih_saat"][:10]
        if gun[:7] == ay.strftime("%Y-%m"):
            kayitlar.setdefault(gun, []).append({
                "tur": "etkinlik",
                "baslik": b["ad"],
                "alt": b["kulup_isim"],
                "saat": "" if b.get("saat_belirsiz") else b["tarih_saat"][11:16],
                "link": url_for("basvuru_detay", basvuru_id=b["id"]),
            })
    donem_adlari = {d["id"]: d["ad"] for d in veri.donemleri_getir()}
    for e in veri.zorunlu_evraklari_getir():
        if e["teslim_tarihi"][:7] == ay.strftime("%Y-%m"):
            kayitlar.setdefault(e["teslim_tarihi"], []).append({
                "tur": "evrak",
                "baslik": f"Son teslim: {e['tur']}",
                "alt": donem_adlari.get(e["donem_id"], ""),
                "saat": "",
                "link": url_for("evrak_takibi", donem=e["donem_id"]),
            })
    for liste in kayitlar.values():
        liste.sort(key=lambda k: (k["tur"] != "evrak", k["saat"]))

    haftalar = calendar.Calendar(firstweekday=0).monthdatescalendar(ay.year, ay.month)
    return render_template(
        "takvim.html",
        ay=ay, ay_adi=f"{AY_ADLARI[ay.month - 1]} {ay.year}",
        onceki=onceki.strftime("%Y-%m"), sonraki=sonraki.strftime("%Y-%m"),
        bu_ay=date.today().strftime("%Y-%m"), bugun=date.today(),
        haftalar=haftalar, kayitlar=kayitlar,
        sayilar={t: sum(1 for l in kayitlar.values() for k in l if k["tur"] == t) for t in ("etkinlik", "evrak")},
    )


# ---------------------------------------------------------------------------
# Belgeler
# ---------------------------------------------------------------------------

def _belge_dosyasini_sil(belge):
    yol = os.path.join(BELGE_KLASORU, belge["dosya"])
    if os.path.exists(yol):
        os.remove(yol)


@app.route("/belgeler", methods=["GET", "POST"])
def belgeler():
    kulupler = _kullanicinin_kulupleri()
    kulup_isimleri = {k["id"]: k["isim"] for k in kulupler}

    if request.method == "POST":
        kulup_id = request.form.get("kulup_id", "")
        tur = request.form.get("tur", "")
        aciklama = request.form.get("aciklama", "").strip()
        donem_id = request.form.get("donem_id", "")
        dosya = request.files.get("dosya")

        if kulup_id not in kulup_isimleri:
            flash("Lütfen bir kulüp seçin.", "hata")
        elif veri.donem_getir(donem_id) is None:
            flash("Lütfen bir dönem seçin.", "hata")
        elif tur not in BELGE_TURLERI:
            flash("Lütfen belge türünü seçin.", "hata")
        elif not dosya or not dosya.filename:
            flash("Lütfen yüklenecek dosyayı seçin.", "hata")
        elif _uzanti(dosya.filename) not in BELGE_UZANTILARI:
            flash("Bu dosya türü desteklenmiyor.", "hata")
        else:
            os.makedirs(BELGE_KLASORU, exist_ok=True)
            belge_id = uuid4().hex
            kayit_adi = f"{belge_id}.{_uzanti(dosya.filename)}"
            dosya.save(os.path.join(BELGE_KLASORU, kayit_adi))
            _islem("Ekleme", "belge", f"{tur}: {secure_filename(dosya.filename) or kayit_adi}", belge_id, kulup_id,
                   aciklama=f"Dönem: {veri.donem_getir(donem_id)['ad']}")
            veri.belge_ekle({
                "id": belge_id,
                "kulup_id": kulup_id,
                "donem_id": donem_id,
                "tur": tur,
                "aciklama": aciklama,
                "orijinal_ad": secure_filename(dosya.filename) or kayit_adi,
                "dosya": kayit_adi,
                "boyut": os.path.getsize(os.path.join(BELGE_KLASORU, kayit_adi)),
                "yukleyen": session.get("ad_soyad", ""),
                "yukleme_tarihi": datetime.now().strftime("%Y-%m-%d %H:%M"),
            })
            flash("Belge yüklendi.", "basarili")
        return redirect(url_for("belgeler", kulup=kulup_id, donem=donem_id or None))

    secili = _baskan_kulubu() or request.args.get("kulup", "")
    donemler = veri.donemleri_getir()
    donem_isimleri = {d["id"]: d["ad"] for d in donemler}
    aktif = next((d["id"] for d in donemler if d["aktif"]), "")
    # Varsayılan filtre aktif dönemdir; "tum" tüm dönemleri gösterir.
    secili_donem = request.args.get("donem", aktif)
    liste = veri.belgeleri_getir(secili or None, None if secili_donem == "tum" else secili_donem or None)
    for belge in liste:
        belge["kulup_isim"] = kulup_isimleri.get(belge["kulup_id"], "—")
        belge["donem_ad"] = donem_isimleri.get(belge.get("donem_id"), "—")
    return render_template(
        "belgeler.html",
        belgeler=liste,
        kulupler=kulupler,
        donemler=donemler,
        aktif_donem_id=aktif,
        secili_donem=secili_donem,
        secili_kulup=secili,
        # Evrak takibinden "yükle" ile gelindiyse yükleme formu hazır doldurulur.
        yukleme_donemi=secili_donem if secili_donem in donem_isimleri else aktif,
        secili_tur=request.args.get("tur", ""),
        belge_turleri=BELGE_TURLERI,
        uzantilar=", ".join(sorted(BELGE_UZANTILARI)),
    )


@app.route("/belgeler/<belge_id>/indir")
def belge_indir(belge_id):
    belge = veri.belge_getir(belge_id)
    if belge is None:
        abort(404)
    return send_from_directory(BELGE_KLASORU, belge["dosya"], as_attachment=True,
                               download_name=belge["orijinal_ad"])


@app.route("/belgeler/<belge_id>/sil", methods=["POST"])
def belge_sil(belge_id):
    """Belgeyi silmez, arşive alır."""
    belge = veri.belge_getir(belge_id)
    if belge is None:
        abort(404)
    veri.belge_arsiv_ayarla(belge_id, session.get("ad_soyad", ""))
    _islem("Arşive alma", "belge", f"{belge['tur']}: {belge['orijinal_ad']}", belge_id, belge["kulup_id"])
    flash("Belge arşive alındı.", "basarili")
    return redirect(_guvenli_yonlendirme(request.form.get("geri")) if request.form.get("geri")
                    else url_for("belgeler"))


# ---------------------------------------------------------------------------
# Arşiv
# ---------------------------------------------------------------------------

ARSIV_TURLERI = {"kulup": "Kulüp", "belge": "Belge", "basvuru": "Etkinlik başvurusu"}


@app.route("/arsiv")
def arsiv():
    kulup_isimleri = {k["id"]: k["isim"] for k in veri.kulupleri_getir(arsiv=None)}
    belgeler = veri.belgeleri_getir(arsiv=True)
    for belge in belgeler:
        belge["kulup_isim"] = kulup_isimleri.get(belge["kulup_id"], "—")
    return render_template("arsiv.html", kulupler=veri.kulupleri_getir(arsiv=True), belgeler=belgeler,
                           basvurular=veri.basvurulari_getir(arsiv=True))


def _arsiv_kaydi_veya_404(tur, kayit_id):
    """(kayıt, görünen ad, kulüp id) döndürür."""
    if tur == "kulup":
        kayit = next((k for k in veri.kulupleri_getir(arsiv=True) if k["id"] == kayit_id), None)
        return (kayit, kayit["isim"], kayit_id) if kayit else abort(404)
    if tur == "belge":
        kayit = veri.belge_getir(kayit_id, arsiv=True)
        return (kayit, f"{kayit['tur']}: {kayit['orijinal_ad']}", kayit["kulup_id"]) if kayit else abort(404)
    if tur == "basvuru":
        kayit = veri.basvuru_getir(kayit_id, arsiv=True)
        return (kayit, kayit["ad"], kayit["kulup_id"]) if kayit else abort(404)
    abort(404)


@app.route("/arsiv/<tur>/<kayit_id>/geri-getir", methods=["POST"])
def arsiv_geri_getir(tur, kayit_id):
    kayit, ad, kulup_id = _arsiv_kaydi_veya_404(tur, kayit_id)
    {"kulup": veri.kulup_arsiv_ayarla, "belge": veri.belge_arsiv_ayarla,
     "basvuru": veri.basvuru_arsiv_ayarla}[tur](kayit_id, None)
    _islem("Arşivden geri getirme", tur, ad, kayit_id, kulup_id)
    if tur != "kulup" and kulup_id not in {k["id"] for k in veri.kulupleri_getir()}:
        flash(f"{ad} geri getirildi, ancak kulübü arşivde olduğu için kulüp geri gelene kadar görünmez.", "basarili")
    else:
        flash(f"{ad} arşivden geri getirildi.", "basarili")
    return redirect(url_for("arsiv"))


@app.route("/arsiv/<tur>/<kayit_id>/kalici-sil", methods=["POST"])
def arsiv_kalici_sil(tur, kayit_id):
    kayit, ad, kulup_id = _arsiv_kaydi_veya_404(tur, kayit_id)
    if tur == "kulup":
        belgeler, ek_dosyalar = veri.kulup_sil(kayit_id)
        for belge in belgeler:
            _belge_dosyasini_sil(belge)
        for dosya in ek_dosyalar:
            _basvuru_dosyasini_sil(dosya)
    elif tur == "belge":
        veri.belge_sil(kayit_id)
        _belge_dosyasini_sil(kayit)
    else:
        for dosya in veri.basvuru_sil(kayit_id):
            _basvuru_dosyasini_sil(dosya)
    _islem("Kalıcı silme", tur, ad, kayit_id, kulup_id)
    flash(f"{ad} kalıcı olarak silindi.", "basarili")
    return redirect(url_for("arsiv"))


# ---------------------------------------------------------------------------
# Kullanıcılar (sadece dekan)
# ---------------------------------------------------------------------------

def _panel_kullanicisi_veya_404(kullanici_adi):
    kullanici = next((k for k in veri.panel_kullanicilari_getir() if k["kullanici_adi"] == kullanici_adi), None)
    if kullanici is None:
        abort(404)
    return kullanici


def _son_aktif_dekan_mi(kullanici_adi):
    dekanlar = [k for k in veri.panel_kullanicilari_getir() if k["rol"] == "dekan" and k["aktif"]]
    return len(dekanlar) == 1 and dekanlar[0]["kullanici_adi"] == kullanici_adi


def _gecici_sifre():
    return secrets.token_urlsafe(9)


@app.route("/kullanicilar")
def kullanicilar():
    kulup_isimleri = {k["id"]: k["isim"] for k in veri.kulupleri_getir(arsiv=None)}
    liste = veri.panel_kullanicilari_getir()
    for k in liste:
        k["kulup_isim"] = kulup_isimleri.get(k.get("kulup_id", ""), "")
    return render_template("kullanicilar.html", kullanicilar=liste, ayarlar=veri.ayarlari_getir())


@app.route("/kullanicilar/yeni", methods=["GET", "POST"])
@app.route("/kullanicilar/<kullanici_adi>/duzenle", methods=["GET", "POST"])
def kullanici_form(kullanici_adi=None):
    mevcut = _panel_kullanicisi_veya_404(kullanici_adi) if kullanici_adi else None
    hata = None
    if request.method == "POST":
        f = request.form
        form = {
            "kullanici_adi": mevcut["kullanici_adi"] if mevcut else f.get("kullanici_adi", "").strip(),
            "ad_soyad": f.get("ad_soyad", "").strip(),
            "eposta": f.get("eposta", "").strip(),
            "rol": f.get("rol", ""),
            "kulup_id": f.get("kulup_id", "") if f.get("rol") == "baskan" else "",
        }
        if not mevcut and not re.fullmatch(r"[A-Za-z0-9._-]{3,32}", form["kullanici_adi"]):
            hata = "Kullanıcı adı 3-32 karakter olmalı; harf, rakam, nokta, alt çizgi ve tire içerebilir."
        elif not mevcut and veri.kullanici_bul(form["kullanici_adi"]):
            hata = "Bu kullanıcı adı zaten kullanılıyor."
        elif not form["ad_soyad"]:
            hata = "Ad soyad boş bırakılamaz."
        elif form["rol"] not in veri.ROLLER:
            hata = "Geçerli bir rol seçin."
        elif form["rol"] == "baskan" and veri.kulup_getir(form["kulup_id"]) is None:
            hata = "Kulüp başkanı için bir kulüp seçin."
        elif mevcut and mevcut["rol"] == "dekan" and form["rol"] != "dekan" and _son_aktif_dekan_mi(mevcut["kullanici_adi"]):
            hata = "Sistemdeki son aktif dekanın rolü değiştirilemez."
        elif not mevcut and len(f.get("sifre", "")) < 8:
            hata = "Şifre en az 8 karakter olmalı."
        else:
            alanlar = {a: form[a] for a in ("ad_soyad", "eposta", "rol", "kulup_id")}
            if not mevcut:
                alanlar.update({"sifre_hash": generate_password_hash(f["sifre"]), "aktif": True,
                                "olusturma": datetime.now().strftime("%Y-%m-%d %H:%M")})
            veri.kullanici_kaydet(form["kullanici_adi"], alanlar)
            _islem("Güncelleme" if mevcut else "Ekleme", "kullanici", f"{form['ad_soyad']} ({form['kullanici_adi']})",
                   form["kullanici_adi"], form["kulup_id"], eski=mevcut or {}, yeni=form)
            flash(f"{form['ad_soyad']} {'güncellendi' if mevcut else 'eklendi'}.", "basarili")
            return redirect(url_for("kullanicilar"))
        mevcut = {**(mevcut or {}), **form}
    return render_template("kullanici_form.html", kullanici=mevcut, duzenleme=kullanici_adi is not None,
                           hata=hata, roller=veri.ROLLER, kulupler=veri.kulupleri_getir(),
                           baskan_girisi=veri.ayarlari_getir()["baskan_girisi"])


@app.route("/kullanicilar/<kullanici_adi>/durum", methods=["POST"])
def kullanici_durum(kullanici_adi):
    kullanici = _panel_kullanicisi_veya_404(kullanici_adi)
    yeni = not kullanici["aktif"]
    if not yeni and kullanici["rol"] == "dekan" and _son_aktif_dekan_mi(kullanici_adi):
        flash("Sistemdeki son aktif dekan pasif yapılamaz.", "hata")
    elif not yeni and kullanici_adi == session.get("kullanici_adi"):
        flash("Kendi hesabınızı pasif yapamazsınız.", "hata")
    else:
        veri.kullanici_kaydet(kullanici_adi, {"aktif": yeni})
        ad = f"{kullanici.get('ad_soyad') or kullanici_adi} ({kullanici_adi})"
        _islem("Güncelleme", "kullanici", ad, kullanici_adi, eski={"aktif": kullanici["aktif"]}, yeni={"aktif": yeni})
        flash(f"{ad} {'aktif' if yeni else 'pasif'} yapıldı.", "basarili")
    return redirect(url_for("kullanicilar"))


@app.route("/kullanicilar/<kullanici_adi>/sifre-sifirla", methods=["POST"])
def kullanici_sifre_sifirla(kullanici_adi):
    kullanici = _panel_kullanicisi_veya_404(kullanici_adi)
    sifre = _gecici_sifre()
    veri.kullanici_kaydet(kullanici_adi, {"sifre_hash": generate_password_hash(sifre)})
    _islem("Şifre sıfırlama", "kullanici", f"{kullanici.get('ad_soyad') or kullanici_adi} ({kullanici_adi})",
           kullanici_adi)
    # Geçici şifre sadece bir kez, bu bildirimde gösterilir; hiçbir yerde düz metin saklanmaz.
    flash(f"{kullanici_adi} için geçici şifre: {sifre} — Bu şifreyi kullanıcıya iletin; "
          f"giriş yaptıktan sonra Hesabım sayfasından değiştirmesini isteyin.", "basarili")
    return redirect(url_for("kullanicilar"))


@app.route("/hesabim", methods=["GET", "POST"])
def hesabim():
    """Giriş yapan kullanıcının kendi şifresini değiştirmesi."""
    kullanici = veri.kullanici_bul(session["kullanici_adi"])
    hata = None
    if request.method == "POST":
        mevcut_sifre = request.form.get("mevcut_sifre", "")
        yeni = request.form.get("yeni_sifre", "")
        if not _sifre_dogru(kullanici, mevcut_sifre):
            hata = "Mevcut şifreniz hatalı."
        elif len(yeni) < 8:
            hata = "Yeni şifre en az 8 karakter olmalı."
        elif yeni != request.form.get("yeni_sifre_tekrar", ""):
            hata = "Yeni şifreler birbiriyle uyuşmuyor."
        else:
            veri.kullanici_kaydet(kullanici["kullanici_adi"], {"sifre_hash": generate_password_hash(yeni)})
            _islem("Şifre değiştirme", "kullanici", f"{session.get('ad_soyad')} ({kullanici['kullanici_adi']})",
                   kullanici["kullanici_adi"])
            flash("Şifreniz değiştirildi.", "basarili")
            return redirect(url_for("hesabim"))
    kulup = veri.kulup_getir(kullanici.get("kulup_id", "")) if kullanici.get("kulup_id") else None
    return render_template("hesabim.html", kullanici=kullanici, kulup=kulup, hata=hata)


# ---------------------------------------------------------------------------
# Ayarlar (sadece dekan)
# ---------------------------------------------------------------------------

@app.route("/ayarlar", methods=["GET", "POST"])
def ayarlar():
    hata = None
    eski = veri.ayarlari_getir()
    if request.method == "POST":
        try:
            sure = int(request.form.get("oturum_suresi_dk", ""))
        except ValueError:
            sure = 0
        if not 5 <= sure <= 480:
            hata = "Oturum süresi 5 ile 480 dakika arasında olmalı."
        else:
            yeni = {"baskan_girisi": request.form.get("baskan_girisi") == "1", "oturum_suresi_dk": sure}
            veri.ayarlari_kaydet(yeni)
            _islem("Güncelleme", "ayar", "Panel ayarları", eski=eski, yeni=yeni)
            flash("Ayarlar kaydedildi.", "basarili")
            return redirect(url_for("ayarlar"))
    baskan_hesaplari = sum(1 for k in veri.panel_kullanicilari_getir() if k["rol"] == "baskan")
    return render_template("ayarlar.html", ayarlar=eski, hata=hata, baskan_hesaplari=baskan_hesaplari)


# ---------------------------------------------------------------------------
# İşlem geçmişi
# ---------------------------------------------------------------------------

@app.route("/islem-gecmisi")
def islem_gecmisi():
    filtre = {a: request.args.get(a, "") for a in ("kullanici_adi", "kulup_id", "kayit_turu", "baslangic", "bitis")}
    islemler = veri.islemleri_getir(**filtre)
    sayfa_boyu = 100
    try:
        sayfa = max(1, int(request.args.get("sayfa", "1")))
    except ValueError:
        sayfa = 1
    toplam_sayfa = max(1, -(-len(islemler) // sayfa_boyu))
    sayfa = min(sayfa, toplam_sayfa)
    kulup_isimleri = {k["id"]: k["isim"] for k in veri.kulupleri_getir(arsiv=None)}
    gosterilen = islemler[(sayfa - 1) * sayfa_boyu: sayfa * sayfa_boyu]
    for i in gosterilen:
        i["kulup_isim"] = kulup_isimleri.get(i.get("kulup_id", ""), "")
        i["tur_ad"] = KAYIT_TURLERI.get(i.get("kayit_turu"), i.get("kayit_turu", ""))
    # Filtre listesindeki kullanıcılar: işlem yapmış herkes (pasif/silinmiş olsa da).
    kullanicilar = sorted({(i.get("kullanici_adi", ""), i.get("kullanici", "")) for i in veri.islemleri_getir()},
                          key=lambda k: k[1].lower())
    return render_template("islem_gecmisi.html", islemler=gosterilen, toplam=len(islemler), sayfa=sayfa,
                           toplam_sayfa=toplam_sayfa, filtre=filtre, kullanicilar=kullanicilar,
                           kulupler=veri.kulupleri_getir(arsiv=None), turler=KAYIT_TURLERI)


# ---------------------------------------------------------------------------
# Duyurular
# ---------------------------------------------------------------------------

HEDEF_TURLERI = {"tum": "Tüm kulüp başkanları", "kategori": "Seçili kategorideki kulüpler", "kulup": "Tek kulüp"}


def _duyuru_hedefi(kaynak):
    """Formdan/sorgudan hedef kulüpleri ve adresleri hesaplar.

    Dönüş: (hedef bilgisi, hata). Hedef bilgisi: tur, deger, ad, kulupler, alicilar, adressiz.
    """
    tur = kaynak.get("hedef_turu", "tum")
    # Formda gizli alan (0) + onay kutusu (1) aynı adla gelir; kutu işaretliyse "1" listededir.
    sadece_aktif = "1" in kaynak.getlist("sadece_aktif")
    danisman_dahil = kaynak.get("danisman_dahil") == "1"
    kulupler = veri.kulupleri_getir()
    if tur == "kategori":
        deger = kaynak.get("kategori", "")
        if not deger:
            return None, "Bir kategori seçin."
        kulupler = [k for k in kulupler if k["kategori"] == deger]
        ad = f"{deger} kategorisindeki kulüpler"
    elif tur == "kulup":
        deger = kaynak.get("kulup_id", "")
        kulupler = [k for k in kulupler if k["id"] == deger]
        if not kulupler:
            return None, "Bir kulüp seçin."
        ad = kulupler[0]["isim"]
        sadece_aktif = False  # tek kulüp seçildiyse durumuna bakılmaz
    else:
        tur, deger, ad = "tum", "", "Tüm kulüp başkanları"
    if sadece_aktif:
        kulupler = [k for k in kulupler if k["durum"] == "aktif"]

    alicilar, adressiz = [], []
    for kulup in kulupler:
        adresler = veri.kulup_eposta_adresleri(kulup, danisman_dahil)
        if adresler:
            alicilar.extend(a for a in adresler if a.lower() not in {x.lower() for x in alicilar})
        else:
            adressiz.append({"id": kulup["id"], "isim": kulup["isim"]})
    return {
        "tur": tur, "deger": deger, "ad": ad, "sadece_aktif": sadece_aktif, "danisman_dahil": danisman_dahil,
        "kulupler": [{"id": k["id"], "isim": k["isim"]} for k in kulupler],
        "alicilar": alicilar, "adressiz": adressiz,
    }, None


def _mailto(duyuru):
    """Duyuruyu e-posta programında açan mailto bağlantısı (alıcılar gizli/BCC)."""
    alan = "bcc" if config.ALICILAR_GIZLI else "to"
    parcalar = [f"{alan}={quote(','.join(duyuru['alicilar']), safe='@,')}",
                f"subject={quote(duyuru['konu'])}", f"body={quote(duyuru['metin'])}"]
    return f"mailto:{quote(config.GONDEREN_ADRES, safe='@')}?" + "&".join(parcalar)


@app.route("/duyurular")
def duyurular():
    return render_template("duyurular.html", duyurular=veri.duyurulari_getir(), hedef_turleri=HEDEF_TURLERI)


@app.route("/duyurular/alicilar")
def duyuru_alicilari():
    """Duyuru formundaki canlı alıcı önizlemesi için JSON."""
    hedef, hata = _duyuru_hedefi(request.args)
    if hata:
        return {"hata": hata}
    return {"kulup_sayisi": len(hedef["kulupler"]), "alici_sayisi": len(hedef["alicilar"]),
            "adressiz": [k["isim"] for k in hedef["adressiz"]]}


@app.route("/duyurular/yeni", methods=["GET", "POST"])
def duyuru_form():
    hata = None
    form = request.form if request.method == "POST" else MultiDict({
        "hedef_turu": request.args.get("hedef_turu", "tum"), "kulup_id": request.args.get("kulup", ""),
        "sadece_aktif": "1"})
    if request.method == "POST":
        hedef, hata = _duyuru_hedefi(request.form)
        konu, metin = request.form.get("konu", "").strip(), request.form.get("metin", "").strip()
        if not hata and (not konu or not metin):
            hata = "Konu ve metin boş bırakılamaz."
        if not hata:
            duyuru = {
                "id": veri.yeni_id(), "tarih": datetime.now().strftime("%Y-%m-%d %H:%M"),
                "gonderen": session.get("ad_soyad", ""), "gonderen_kullanici": session.get("kullanici_adi", ""),
                "konu": konu, "metin": metin, "hedef_turu": hedef["tur"], "hedef_deger": hedef["deger"],
                "hedef_ad": hedef["ad"], "kulupler": hedef["kulupler"], "alicilar": hedef["alicilar"],
                "adressiz": hedef["adressiz"], "danisman_dahil": hedef["danisman_dahil"],
                "yontem": config.EPOSTA_YONTEMI,
            }
            veri.duyuru_ekle(duyuru)
            _islem("Ekleme", "duyuru", konu, duyuru["id"],
                   hedef["deger"] if hedef["tur"] == "kulup" else "",
                   aciklama=f"{hedef['ad']} · {len(hedef['kulupler'])} kulüp, {len(hedef['alicilar'])} alıcı")
            flash("Duyuru arşive kaydedildi. Şimdi e-posta programınızda açıp gönderebilirsiniz.", "basarili")
            return redirect(url_for("duyuru_detay", duyuru_id=duyuru["id"]))
    return render_template("duyuru_form.html", form=form, hata=hata, hedef_turleri=HEDEF_TURLERI,
                           kategoriler=veri.kategorileri_getir(), kulupler=veri.kulupleri_getir())


@app.route("/duyurular/<duyuru_id>")
def duyuru_detay(duyuru_id):
    duyuru = veri.duyuru_getir(duyuru_id)
    if duyuru is None:
        abort(404)
    mailto = _mailto(duyuru) if duyuru["alicilar"] else ""
    return render_template("duyuru_detay.html", duyuru=duyuru, mailto=mailto,
                           mailto_uzun=len(mailto) > config.MAILTO_UZUNLUK_SINIRI,
                           gizli=config.ALICILAR_GIZLI, hedef_turleri=HEDEF_TURLERI)


# ---------------------------------------------------------------------------
# İletişim listesi (Excel)
# ---------------------------------------------------------------------------

def _excel_basligi(sayfa, basliklar, genislikler):
    from openpyxl.styles import Font, PatternFill

    sayfa.append(basliklar)
    for hucre in sayfa[sayfa.max_row]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="1B2F8F")
    for i, genislik in enumerate(genislikler):
        sayfa.column_dimensions[chr(65 + i)].width = genislik
    # Hücreye erişmek o satırı oluşturur; konum metin olarak verilir.
    sayfa.freeze_panes = f"A{sayfa.max_row + 1}"


@app.route("/iletisim-listesi.xlsx")
def iletisim_listesi():
    from openpyxl import Workbook

    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Kulüpler"
    _excel_basligi(sayfa, ["Kulüp", "Kategori", "Durum", "Başkan", "Başkan / Kulüp E-posta", "Danışman",
                           "Danışman E-posta", "Danışman Telefon"], (36, 22, 8, 24, 34, 34, 30, 16))
    for k in veri.kulupleri_getir():
        sayfa.append([k["isim"], k["kategori"], "Aktif" if k["durum"] == "aktif" else "Pasif",
                      k["baskan"] or "Atanmamış", ", ".join(veri.kulup_eposta_adresleri(k)),
                      k["danisman"]["isim"], k["danisman"]["eposta"], k["danisman"]["telefon"]])

    sayfa = kitap.create_sheet("Danışmanlar")
    _excel_basligi(sayfa, ["Unvan", "Ad Soyad", "Fakülte", "Bölüm", "E-posta", "Telefon", "Kulüpler"],
                   (16, 26, 28, 30, 30, 16, 50))
    for d in veri.danismanlari_getir():
        sayfa.append([d["unvan"], d["ad_soyad"], d["fakulte"], d["bolum"], d["eposta"], d["telefon"],
                      ", ".join(k["isim"] for k in d["kulupler"])])

    cikti = io.BytesIO()
    kitap.save(cikti)
    cikti.seek(0)
    return send_file(cikti, as_attachment=True, download_name=f"kulup-iletisim-listesi-{date.today().isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------------------
# Genel arama (üst çubuk)
# ---------------------------------------------------------------------------

def _arama_metni(metin):
    """Türkçe büyük/küçük harf duyarsız karşılaştırma için (İ/I dönüşümü dahil)."""
    return (metin or "").replace("İ", "i").replace("I", "ı").lower()


@app.route("/ara")
def ara():
    """Kulüp, başkan, danışman ve öğrenci (numara veya ad) araması; JSON döner."""
    q = _arama_metni(request.args.get("q", "").strip())
    if len(q) < 2:
        return {"sonuclar": []}
    sonuclar = []
    kulupler = veri.kulupleri_getir()
    for k in kulupler:
        if q in _arama_metni(k["isim"]):
            sonuclar.append({"tur": "Kulüp", "baslik": k["isim"], "alt": k["kategori"],
                             "link": url_for("kulup_detay", kulup_id=k["id"])})
    for k in kulupler:
        if k["baskan"] and q in _arama_metni(k["baskan"]):
            sonuclar.append({"tur": "Başkan", "baslik": k["baskan"], "alt": k["isim"],
                             "link": url_for("kulup_detay", kulup_id=k["id"])})
    for d in veri.danismanlari_getir():
        if q in _arama_metni(d["ad_soyad"]):
            sonuclar.append({"tur": "Danışman", "baslik": d["tam_ad"],
                             "alt": ", ".join(k["isim"] for k in d["kulupler"]) or "Kulübü yok",
                             "link": url_for("danisman_detay", danisman_id=d["id"])})
    # Öğrenciler: numara veya ad soyad; bir öğrenci birden fazla kulüpteyse tek satırda gösterilir.
    ogrenciler = {}
    for u in veri.tum_uyeleri_getir():
        if q in _arama_metni(u["ogrenci_no"]) or (u["ad_soyad"] and q in _arama_metni(u["ad_soyad"])):
            kayit = ogrenciler.setdefault(u["ogrenci_no"], {"uye": u, "kulupler": []})
            kayit["kulupler"].append((u["kulup_id"], u["kulup_isim"]))
    for no, kayit in list(ogrenciler.items())[:15]:
        sonuclar.append({"tur": "Öğrenci", "baslik": f"{kayit['uye']['ad_soyad'] or 'Ad girilmemiş'} · {no}",
                         "alt": ", ".join(isim for _, isim in kayit["kulupler"]),
                         "link": url_for("kulup_detay", kulup_id=kayit["kulupler"][0][0]) + "#uyeler"})
    return {"sonuclar": sonuclar[:40], "fazla": len(sonuclar) > 40}


# ---------------------------------------------------------------------------
# Excel'den üye içe aktarma (önizleme + onay)
# ---------------------------------------------------------------------------

SABLON_SUTUNLARI = ["Ad Soyad", "Öğrenci No", "Bölüm", "Sınıf", "Telefon", "E-posta", "Katılım Tarihi (GG.AA.YYYY)"]
SABLON_ALANLARI = ["ad_soyad", "ogrenci_no", "bolum", "sinif", "telefon", "eposta", "katilim_tarihi"]


@app.route("/kulupler/<kulup_id>/uyeler/sablon.xlsx")
def uye_sablonu(kulup_id):
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    kulup = _kulup_veya_404(kulup_id)
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Üyeler"
    _excel_basligi(sayfa, SABLON_SUTUNLARI, (26, 14, 28, 10, 16, 30, 26))
    sayfa["A1"].comment = Comment(f"{kulup['isim']} üye listesi. Ad Soyad ve Öğrenci No zorunludur. "
                                  "Başlık satırını değiştirmeyin.", "Öğrenci Dekanlığı")
    # Metin olarak saklansın diye öğrenci no ve telefon sütunları metin biçiminde.
    for satir in range(2, 502):
        for sutun in ("B", "E"):
            sayfa[f"{sutun}{satir}"].number_format = "@"
    cikti = io.BytesIO()
    kitap.save(cikti)
    cikti.seek(0)
    return send_file(cikti, as_attachment=True, download_name=f"{kulup_id}-uye-sablonu.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _hucre_metni(deger):
    if deger is None:
        return ""
    if isinstance(deger, float) and deger.is_integer():
        deger = int(deger)
    return str(deger).strip()


def _tarih_coz(deger):
    """Excel tarih hücresi veya 'GG.AA.YYYY' / 'YYYY-AA-GG' metnini 'YYYY-AA-GG' yapar."""
    if deger in (None, ""):
        return "", None
    if isinstance(deger, datetime):
        return deger.date().isoformat(), None
    if isinstance(deger, date):
        return deger.isoformat(), None
    metin = str(deger).strip()
    for bicim in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(metin, bicim).date().isoformat(), None
        except ValueError:
            pass
    return "", f"Katılım tarihi anlaşılamadı: {metin}"


def _excel_uyelerini_oku(dosya, kulup_id):
    """Excel'deki satırları okuyup doğrular; her satıra durum ekler."""
    from openpyxl import load_workbook

    sayfa = load_workbook(dosya, read_only=True, data_only=True).active
    satirlar = list(sayfa.iter_rows(values_only=True))
    if not satirlar:
        raise ValueError("Dosya boş.")
    baslik = [_arama_metni(_hucre_metni(h)) for h in satirlar[0]]
    beklenen = [_arama_metni(b) for b in SABLON_SUTUNLARI]
    if baslik[:2] != beklenen[:2]:
        raise ValueError("Başlık satırı şablonla uyuşmuyor. Lütfen panelden indirilen şablonu kullanın.")

    mevcut = {u["ogrenci_no"]: u for u in veri.uyeleri_getir(kulup_id)}
    tum_ogrenciler = {u["ogrenci_no"]: u for u in veri.tum_uyeleri_getir()}
    gorulen, sonuc = set(), []
    for no, ham in enumerate(satirlar[1:], start=2):
        if not any(_hucre_metni(h) for h in ham):
            continue
        ham = list(ham) + [None] * len(SABLON_ALANLARI)
        uye = {alan: _hucre_metni(ham[i]) for i, alan in enumerate(SABLON_ALANLARI) if alan != "katilim_tarihi"}
        uye["katilim_tarihi"], tarih_hatasi = _tarih_coz(ham[6])
        hatalar = []
        if not uye["ad_soyad"] or not uye["ogrenci_no"]:
            hatalar.append("Ad soyad ve öğrenci no zorunlu.")
        elif not re.fullmatch(r"[0-9A-Za-zçğıöşüÇĞİÖŞÜ]{1,32}", uye["ogrenci_no"]):
            hatalar.append("Öğrenci numarası geçersiz.")
        if uye["eposta"] and not EPOSTA_DESENI.fullmatch(uye["eposta"]):
            hatalar.append("E-posta geçersiz.")
        if tarih_hatasi:
            hatalar.append(tarih_hatasi)
        if uye["ogrenci_no"] in gorulen:
            hatalar.append("Bu öğrenci numarası dosyada birden fazla kez geçiyor.")
        gorulen.add(uye["ogrenci_no"])

        if hatalar:
            durum, aciklama = "hata", " ".join(hatalar)
        elif uye["ogrenci_no"] in mevcut:
            eski = mevcut[uye["ogrenci_no"]]
            farkli = [a for a in SABLON_ALANLARI if a != "ogrenci_no" and uye[a] and uye[a] != eski.get(a, "")]
            durum = "guncelleme" if farkli else "ayni"
            aciklama = ("Zaten üye; güncellenecek: " + ", ".join(ALAN_ADLARI[a] for a in farkli)) if farkli \
                else "Zaten üye; değişiklik yok."
        else:
            durum = "yeni"
            diger = tum_ogrenciler.get(uye["ogrenci_no"])
            aciklama = f"Yeni üye (başka kulüpte de kayıtlı: {diger['kulup_isim']})" if diger else "Yeni üye"
        sonuc.append({"satir": no, "uye": uye, "durum": durum, "aciklama": aciklama})
    return sonuc


@app.route("/kulupler/<kulup_id>/uyeler/ice-aktar", methods=["GET", "POST"])
def uye_ice_aktar(kulup_id):
    kulup = _kulup_veya_404(kulup_id)
    satirlar, anahtar, hata = None, None, None
    if request.method == "POST":
        dosya = request.files.get("dosya")
        if not dosya or not dosya.filename:
            hata = "Lütfen bir Excel dosyası seçin."
        elif _uzanti(dosya.filename) != "xlsx":
            hata = "Sadece .xlsx dosyası yüklenebilir. Panelden indirilen şablonu kullanın."
        else:
            try:
                satirlar = _excel_uyelerini_oku(dosya, kulup_id)
            except Exception as e:  # bozuk/uyumsuz dosya
                hata = str(e) if isinstance(e, ValueError) else "Dosya okunamadı. Geçerli bir Excel (.xlsx) dosyası yükleyin."
            if satirlar is not None and not satirlar:
                hata, satirlar = "Dosyada üye satırı bulunamadı.", None
        if satirlar:
            # Önizleme sunucuda geçici olarak saklanır; onaylanınca okunup silinir.
            os.makedirs(ICE_AKTARIM_KLASORU, exist_ok=True)
            anahtar = secrets.token_hex(16)
            with open(os.path.join(ICE_AKTARIM_KLASORU, f"{anahtar}.json"), "w", encoding="utf-8") as f:
                json.dump({"kulup_id": kulup_id, "kullanici": session.get("kullanici_adi"), "satirlar": satirlar},
                          f, ensure_ascii=False)
    sayilar = {d: sum(1 for s in satirlar or [] if s["durum"] == d) for d in ("yeni", "guncelleme", "ayni", "hata")}
    return render_template("uye_ice_aktar.html", kulup=kulup, satirlar=satirlar, anahtar=anahtar, hata=hata,
                           sayilar=sayilar, sutunlar=SABLON_SUTUNLARI)


@app.route("/kulupler/<kulup_id>/uyeler/ice-aktar/onayla", methods=["POST"])
def uye_ice_aktar_onayla(kulup_id):
    kulup = _kulup_veya_404(kulup_id)
    anahtar = request.form.get("anahtar", "")
    yol = os.path.join(ICE_AKTARIM_KLASORU, f"{anahtar}.json")
    if not re.fullmatch(r"[0-9a-f]{32}", anahtar) or not os.path.exists(yol):
        flash("Önizleme bulunamadı veya süresi doldu. Dosyayı yeniden yükleyin.", "hata")
        return redirect(url_for("uye_ice_aktar", kulup_id=kulup_id))
    with open(yol, encoding="utf-8") as f:
        onizleme = json.load(f)
    os.remove(yol)
    if onizleme["kulup_id"] != kulup_id or onizleme["kullanici"] != session.get("kullanici_adi"):
        abort(403)
    eklenen = guncellenen = 0
    for s in onizleme["satirlar"]:
        if s["durum"] not in ("yeni", "guncelleme"):
            continue
        uye = s["uye"]
        # Güncellemede Excel'de boş bırakılan alanlar mevcut bilgiyi silmez.
        eski = next((u for u in veri.uyeleri_getir(kulup_id) if u["ogrenci_no"] == uye["ogrenci_no"]), {})
        bilgiler = {a: uye.get(a) or eski.get(a, "") for a in veri.OGRENCI_ALANLARI}
        veri.uye_kaydet(kulup_id, uye["ogrenci_no"], bilgiler, uye["katilim_tarihi"] or eski.get("katilim_tarihi", ""))
        eklenen += s["durum"] == "yeni"
        guncellenen += s["durum"] == "guncelleme"
    _islem("İçe aktarma", "uye", f"{kulup['isim']} üye listesi", kulup_id, kulup_id,
           aciklama=f"Excel'den {eklenen} yeni üye eklendi, {guncellenen} üye güncellendi.")
    flash(f"İçe aktarma tamamlandı: {eklenen} yeni üye eklendi, {guncellenen} üye güncellendi.", "basarili")
    return redirect(url_for("kulup_detay", kulup_id=kulup_id) + "#uyeler")


@app.template_filter("boyut")
def boyut_filtresi(bayt):
    for birim in ("B", "KB", "MB"):
        if bayt < 1024:
            return f"{bayt:.0f} {birim}"
        bayt /= 1024
    return f"{bayt:.1f} GB"


@app.template_filter("para")
def para_filtresi(tutar):
    """1234.5 -> '1.234,50 TL'"""
    return f"{tutar:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " TL"


@app.template_filter("tarih")
def tarih_filtresi(deger):
    """'2026-10-24' -> '24.10.2026' (boşsa boş döner)."""
    try:
        return datetime.strptime(deger[:10], "%Y-%m-%d").strftime("%d.%m.%Y") + deger[10:]
    except (TypeError, ValueError):
        return deger or ""


if __name__ == "__main__":
    # debug=True: kod değiştikçe sunucu otomatik yeniden başlar.
    app.run(debug=True)
